from datetime import date
from decimal import Decimal
from io import StringIO
import os
import re
import shutil
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase, override_settings
from django.utils import timezone
from docx import Document
from PIL import Image
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from .models import (
	Brand,
	CarData,
	Configuration,
	Generation,
	Model,
	Protocol,
	ProtocolBrake,
	ProtocolMeasurement,
	ProtocolFormSettings,
	ProtocolPhoto,
	ProtocolStorageSettings,
)
from .serializers import (
	ProtocolCreateSerializer,
	ConfigurationSerializer,
	ProtocolBrakeSerializer,
	ProtocolLightSerializer,
	ProtocolMeasurementSerializer,
	ProtocolPowerSupplySerializer,
	ProtocolTestConditionSerializer,
)
from .services.protocol_docx.calculations import (
	build_uncertainty_values,
	calc_axle_mass_range_kg,
	calc_drom_full_mass_kg,
	count_seats,
	format_mass_kg,
	get_vehicle_mass_for_brake_calc,
	calc_light_absorption_uncertainty,
	calc_axle_load_sum_kg,
	standard_headlight_uncertainty_cd,
	u_control_force_n,
	u_glass_transparency_pct,
	u_noise_db,
	u_axle_load_kg,
	u_speed_kmh,
	u_steering_backlash_deg,
	u_turn_signal_frequency_hz,
	u_vehicle_height_mm,
)
from .services.evaluation_sheet import EVALUATION_ROWS, build_evaluation_context
from .services.protocol_docx_generator import generate_protocol_docx
from .services.xlsx_renderer import render_xlsx
from .services.protocol_storage import (
	photo_file_stem,
	protocol_number_label,
	validate_storage_dirs,
)
from .services.protocol_docx.applicability import build_applicability_values
from .services.protocol_docx.context import (
	STATIC_CONCLUSIONS,
	build_dash_conclusion_values,
	build_protocol_docx_context,
	build_submitted_documents,
	build_light_device_row_values,
	build_dynamic_result_values,
	build_eco_values,
	build_front_fog_values,
	build_rear_fog_values,
	build_tire_depth_values,
)
from .services.protocol_docx.renderer import process_document_conditions, render_protocol_docx
from .services.protocol_docx.placeholder_aliases import (
	add_protocol_placeholder_aliases,
	load_placeholder_aliases,
)


class ProtocolWorkflowTests(TestCase):
	def setUp(self):
		self.client = APIClient()
		self.measurer = self.create_user("measurer", "measurer")
		self.other_measurer = self.create_user("other-measurer", "measurer")
		self.operator = self.create_user("operator", "operator")
		self.manager = self.create_user("manager", "manager")
		self.brand = Brand.objects.create(
			name="Test KIA",
			link="https://example.com/test-kia",
		)
		self.model = Model.objects.create(
			name="Test RIO",
			link="https://example.com/test-rio",
			brand=self.brand,
		)
		self.generation = Generation.objects.create(
			name="Test generation",
			link="https://example.com/test-generation",
			model=self.model,
			date_start="01.2020",
			date_end="12.2021",
		)
		self.other_generation = Generation.objects.create(
			name="Other generation",
			link="https://example.com/other-generation",
			model=self.model,
			date_start="01.2022",
			date_end="12.2023",
		)
		self.configuration = Configuration.objects.create(
			name="Test configuration",
			link="https://example.com/test-configuration",
			generation=self.generation,
			date_start="01.2020",
			date_end="12.2021",
		)
		self.other_configuration = Configuration.objects.create(
			name="Other configuration",
			link="https://example.com/other-configuration",
			generation=self.other_generation,
			date_start="01.2022",
			date_end="12.2023",
		)
		CarData.objects.create(
			configuration=self.configuration,
			manufacture_year=2020,
			front_tires="185/65R15",
			rear_tires="185/65R15",
			fuel_type="petrol",
			transmission="automatic",
			drive_type="front",
			vehicle_length_mm=3800,
			vehicle_width_mm=1680,
			vehicle_height_mm=1510,
			vehicle_weight_kg=990,
			engine_model="Test engine",
			engine_power_kw=45,
			cylinder_layout="inline",
			cylinders_count=3,
			turbo_present=False,
			front_brakes="дисковые",
			rear_brakes="дисковые",
		)
		CarData.objects.create(
			configuration=self.other_configuration,
			manufacture_year=2022,
			front_tires="195/55R16",
			rear_tires="195/55R16",
			fuel_type="petrol",
			transmission="automatic",
			drive_type="front",
		)
		self.protocol = Protocol.objects.create(
			user=self.measurer,
			model=self.model,
			protocol_date=date(2020, 5, 1),
			protocol_number="00001",
			owner_name="Тестовый владелец",
			brand_name="KIA",
			commercial_name="RIO",
			status="measurement",
		)

	@staticmethod
	def create_user(username, role):
		user = User.objects.create_user(username=username, password="test-password")
		user.groups.add(Group.objects.get_or_create(name=role)[0])
		return user

	def authenticate(self, user):
		self.client.force_authenticate(user=user)

	def post(self, path, data=None):
		return self.client.post(path, data or {}, format="json")

	def test_full_protocol_workflow_reaches_approved(self):
		self.authenticate(self.measurer)

		start_response = self.post(
			f"/cars/protocols/{self.protocol.id}/start-editing/"
		)
		self.assertEqual(start_response.status_code, 200)

		submit_response = self.post(
			f"/cars/protocols/{self.protocol.id}/submit-to-operator/"
		)
		self.assertEqual(submit_response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "operator")
		self.assertIsNone(self.protocol.locked_by_id)

		self.authenticate(self.operator)
		self.assertEqual(
			self.post(f"/cars/protocols/{self.protocol.id}/start-editing/").status_code,
			200,
		)
		review_response = self.post(
			f"/cars/protocols/{self.protocol.id}/submit-for-review/"
		)
		self.assertEqual(review_response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "review")
		self.assertIsNone(self.protocol.locked_by_id)

		self.authenticate(self.manager)
		approve_response = self.post(
			f"/cars/protocols/{self.protocol.id}/approve/"
		)
		self.assertEqual(approve_response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "approved")
		self.assertIsNone(self.protocol.locked_by_id)

	def test_reviewer_can_return_protocol_for_revision(self):
		self.protocol.status = "review"
		self.protocol.save(update_fields=["status"])
		self.authenticate(self.manager)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/cancel/",
			{"revision_comment": "Нужно уточнить результаты замера"},
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "revision")
		self.assertTrue(self.protocol.returned_for_revision)
		self.assertEqual(
			self.protocol.revision_comment,
			"Нужно уточнить результаты замера",
		)
		self.assertEqual(self.protocol.cancelled_by_id, self.manager.id)
		self.assertIsNone(self.protocol.locked_by_id)

	def test_measurer_cannot_access_another_measurers_protocol(self):
		self.authenticate(self.other_measurer)

		response = self.client.get(f"/cars/protocols/{self.protocol.id}/")

		self.assertEqual(response.status_code, 403)

	def test_locked_protocol_cannot_be_opened_by_second_user(self):
		self.authenticate(self.measurer)
		self.assertEqual(
			self.post(f"/cars/protocols/{self.protocol.id}/start-editing/").status_code,
			200,
		)

		self.authenticate(self.operator)
		response = self.post(
			f"/cars/protocols/{self.protocol.id}/start-editing/"
		)

		self.assertEqual(response.status_code, 423)
		self.assertEqual(response.data["locked_by_id"], self.measurer.id)

	def test_superuser_can_lock_and_update_protocol(self):
		admin = User.objects.create_superuser(
			username="admin",
			email="admin@example.com",
			password="test-password",
		)
		self.authenticate(admin)

		lock_response = self.post(
			f"/cars/protocols/{self.protocol.id}/start-editing/"
		)
		self.assertEqual(lock_response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.locked_by_id, admin.id)

		update_response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"owner_name": "Updated by admin"},
			format="json",
		)
		self.assertEqual(update_response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.owner_name, "Updated by admin")

	def test_superuser_can_update_approved_protocol_despite_foreign_lock(self):
		self.protocol.status = "approved"
		self.protocol.locked_by = self.other_measurer
		self.protocol.locked_at = timezone.now()
		self.protocol.save(update_fields=["status", "locked_by", "locked_at"])
		admin = User.objects.create_superuser(
			username="admin-all-rights",
			email="admin-all-rights@example.com",
			password="test-password",
		)
		self.authenticate(admin)

		response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"owner_name": "Updated by full-access admin"},
			format="json",
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.owner_name, "Updated by full-access admin")

	def test_measurement_api_saves_pneumatic_suspension_presence(self):
		ProtocolMeasurement.objects.create(protocol=self.protocol)
		self.authenticate(self.measurer)
		self.assertEqual(
			self.post(f"/cars/protocols/{self.protocol.id}/start-editing/").status_code,
			200,
		)

		response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/measurement/update/",
			{"pneumatic_suspension_present": True},
			format="json",
		)

		self.assertEqual(response.status_code, 200)
		measurement = ProtocolMeasurement.objects.get(protocol=self.protocol)
		self.assertTrue(measurement.pneumatic_suspension_present)

	def test_operator_cannot_approve_protocol(self):
		self.protocol.status = "review"
		self.protocol.save(update_fields=["status"])
		self.authenticate(self.operator)

		response = self.post(f"/cars/protocols/{self.protocol.id}/approve/")

		self.assertEqual(response.status_code, 403)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "review")

	def test_revision_protocol_can_be_sent_back_to_review(self):
		self.protocol.status = "revision"
		self.protocol.returned_for_revision = True
		self.protocol.revision_comment = "Исправить данные"
		self.protocol.save(
			update_fields=["status", "returned_for_revision", "revision_comment"]
		)

		self.authenticate(self.operator)
		self.assertEqual(
			self.post(f"/cars/protocols/{self.protocol.id}/start-editing/").status_code,
			200,
		)
		response = self.post(
			f"/cars/protocols/{self.protocol.id}/submit-for-review/"
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "review")
		self.assertFalse(self.protocol.returned_for_revision)
		self.assertIsNone(self.protocol.revision_comment)

	def test_only_lock_owner_can_send_heartbeat(self):
		self.authenticate(self.measurer)
		self.assertEqual(
			self.post(f"/cars/protocols/{self.protocol.id}/start-editing/").status_code,
			200,
		)

		self.authenticate(self.operator)
		response = self.post(f"/cars/protocols/{self.protocol.id}/heartbeat/")

		self.assertEqual(response.status_code, 403)

	def test_reviewer_can_release_another_users_lock(self):
		self.authenticate(self.measurer)
		self.assertEqual(
			self.post(f"/cars/protocols/{self.protocol.id}/start-editing/").status_code,
			200,
		)

		self.authenticate(self.manager)
		response = self.post(
			f"/cars/protocols/{self.protocol.id}/manager-release-lock/"
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertIsNone(self.protocol.locked_by_id)
		self.assertIsNone(self.protocol.locked_at)

	def test_superuser_returns_approved_protocol_to_operator_and_can_edit(self):
		self.protocol.status = "approved"
		self.protocol.returned_for_revision = True
		self.protocol.revision_comment = "Previously returned"
		self.protocol.cancelled_by = self.manager
		self.protocol.cancelled_at = timezone.now()
		self.protocol.save(
			update_fields=[
				"status",
				"returned_for_revision",
				"revision_comment",
				"cancelled_by",
				"cancelled_at",
			]
		)
		admin = User.objects.create_superuser(
			username="admin-return",
			email="admin-return@example.com",
			password="test-password",
		)
		self.authenticate(admin)

		response = self.post(f"/cars/protocols/{self.protocol.id}/return-to-draft/")

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "operator")
		self.assertEqual(self.protocol.locked_by_id, admin.id)
		self.assertIsNotNone(self.protocol.locked_at)
		self.assertFalse(self.protocol.returned_for_revision)
		self.assertIsNone(self.protocol.revision_comment)
		self.assertIsNone(self.protocol.cancelled_by_id)
		self.assertIsNone(self.protocol.cancelled_at)

		update_response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"owner_name": "Reopened by admin"},
			format="json",
		)
		self.assertEqual(update_response.status_code, 200)

	def test_operator_cannot_return_approved_protocol_to_operator(self):
		self.protocol.status = "approved"
		self.protocol.save(update_fields=["status"])
		self.authenticate(self.operator)

		response = self.post(f"/cars/protocols/{self.protocol.id}/return-to-draft/")

		self.assertEqual(response.status_code, 403)

	def test_operator_list_excludes_measurement_and_review_protocols(self):
		Protocol.objects.create(
			user=self.measurer,
			protocol_number="00002",
			status="operator",
		)
		Protocol.objects.create(
			user=self.measurer,
			protocol_number="00003",
			status="review",
		)

		self.authenticate(self.operator)
		response = self.client.get("/cars/protocols/")

		self.assertEqual(response.status_code, 200)
		self.assertEqual(
			{item["status"] for item in response.data["results"]},
			{"operator"},
		)

	def test_approved_protocol_cannot_be_started_for_editing(self):
		self.protocol.status = "approved"
		self.protocol.save(update_fields=["status"])
		self.authenticate(self.measurer)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/start-editing/"
		)

		self.assertEqual(response.status_code, 403)

	def test_configuration_must_belong_to_selected_generation(self):
		self.protocol.locked_by = self.measurer
		self.protocol.locked_at = timezone.now()
		self.protocol.save(update_fields=["locked_by", "locked_at"])
		self.authenticate(self.measurer)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/select-configuration/",
			{
				"generation_id": self.generation.id,
				"configuration_id": self.other_configuration.id,
				"manufacture_date": "2020-05",
			},
		)

		self.assertEqual(response.status_code, 400)
		self.protocol.refresh_from_db()
		self.assertIsNone(self.protocol.configuration_id)

	def test_configuration_must_match_manufacture_month(self):
		self.protocol.locked_by = self.measurer
		self.protocol.locked_at = timezone.now()
		self.protocol.save(update_fields=["locked_by", "locked_at"])
		self.authenticate(self.measurer)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/select-configuration/",
			{
				"generation_id": self.generation.id,
				"configuration_id": self.configuration.id,
				"manufacture_date": "2022-05",
			},
		)

		self.assertEqual(response.status_code, 400)
		self.protocol.refresh_from_db()
		self.assertIsNone(self.protocol.configuration_id)

	def test_valid_generation_configuration_and_month_are_saved(self):
		self.protocol.locked_by = self.measurer
		self.protocol.locked_at = timezone.now()
		self.protocol.save(update_fields=["locked_by", "locked_at"])
		self.authenticate(self.measurer)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/select-configuration/",
			{
				"generation_id": self.generation.id,
				"configuration_id": self.configuration.id,
				"manufacture_date": "2020-05",
			},
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.generation_id, self.generation.id)
		self.assertEqual(self.protocol.configuration_id, self.configuration.id)
		self.assertEqual(self.protocol.manufacture_date, date(2020, 5, 1))
		measurement = ProtocolMeasurement.objects.get(protocol=self.protocol)
		self.assertEqual(measurement.vehicle_length_mm, Decimal("3800"))
		self.assertEqual(measurement.vehicle_width_mm, Decimal("1680"))
		self.assertEqual(measurement.vehicle_height_mm, Decimal("1510"))
		self.assertEqual(measurement.vehicle_weight_kg, Decimal("990"))
		self.assertEqual(measurement.wheel_formula, "4x2_front")
		self.assertEqual(measurement.engine_model, "Test engine")
		self.assertEqual(measurement.fuel_type, "petrol")
		self.assertEqual(
			ProtocolBrake.objects.get(protocol=self.protocol).service_brake_type,
			"disc_disc",
		)
		measurement = ProtocolMeasurement.objects.get(protocol=self.protocol)
		self.assertEqual(measurement.vehicle_length_mm, Decimal("3800"))
		self.assertEqual(measurement.vehicle_width_mm, Decimal("1680"))
		self.assertEqual(measurement.vehicle_height_mm, Decimal("1510"))
		self.assertEqual(measurement.vehicle_weight_kg, Decimal("990"))
		self.assertEqual(measurement.wheel_formula, "4x2_front")
		self.assertEqual(measurement.engine_model, "Test engine")
		self.assertEqual(measurement.fuel_type, "petrol")

	def test_valid_generation_configuration_can_be_saved_without_manufacture_month(self):
		self.protocol.locked_by = self.measurer
		self.protocol.locked_at = timezone.now()
		self.protocol.save(update_fields=["locked_by", "locked_at"])
		self.authenticate(self.measurer)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/select-configuration/",
			{
				"generation_id": self.generation.id,
				"configuration_id": self.configuration.id,
			},
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.configuration_id, self.configuration.id)
		self.assertIsNone(self.protocol.manufacture_date)

	def test_configuration_import_normalizes_full_drive_to_four_by_four(self):
		CarData.objects.filter(configuration=self.configuration).update(drive_type="full")
		self.protocol.locked_by = self.measurer
		self.protocol.locked_at = timezone.now()
		self.protocol.save(update_fields=["locked_by", "locked_at"])
		self.authenticate(self.measurer)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/select-configuration/",
			{
				"generation_id": self.generation.id,
				"configuration_id": self.configuration.id,
				"manufacture_date": "2020-05",
			},
		)

		self.assertEqual(response.status_code, 200)
		measurement = ProtocolMeasurement.objects.get(protocol=self.protocol)
		self.assertEqual(measurement.wheel_formula, "4x4")

	def test_configuration_payload_exposes_car_data_dimensions(self):
		payload = ConfigurationSerializer(self.configuration).data

		self.assertEqual(payload["vehicle_length_mm"], 3800)
		self.assertEqual(payload["vehicle_width_mm"], 1680)
		self.assertEqual(payload["vehicle_height_mm"], 1510)
		self.assertEqual(payload["vehicle_weight_kg"], 990)
		self.assertEqual(payload["engine_model"], "Test engine")
		self.assertEqual(payload["engine_power_kw"], 45)

	def test_configuration_import_preserves_measured_and_dash_values(self):
		measurement = ProtocolMeasurement.objects.create(
			protocol=self.protocol,
			vehicle_length_mm=Decimal("4001"),
			vehicle_width_mm=None,
			dash_fields=["vehicle_width_mm"],
		)
		self.protocol.locked_by = self.measurer
		self.protocol.locked_at = timezone.now()
		self.protocol.save(update_fields=["locked_by", "locked_at"])
		self.authenticate(self.measurer)

		response = self.post(
			f"/cars/protocols/{self.protocol.id}/select-configuration/",
			{
				"generation_id": self.generation.id,
				"configuration_id": self.configuration.id,
				"manufacture_date": "2020-05",
			},
		)

		self.assertEqual(response.status_code, 200)
		measurement.refresh_from_db()
		self.assertEqual(measurement.vehicle_length_mm, Decimal("4001"))
		self.assertIsNone(measurement.vehicle_width_mm)
		self.assertEqual(measurement.vehicle_height_mm, Decimal("1510"))

	def test_measurer_cannot_access_foreign_protocol_resources(self):
		photo = ProtocolPhoto.objects.create(
			protocol=self.protocol,
			photo_type="odometer_photo",
			file_path="e2e/foreign-photo.jpg",
		)
		self.authenticate(self.other_measurer)

		protected_get_urls = [
			f"/cars/protocols/{self.protocol.id}/",
			f"/cars/protocols/{self.protocol.id}/full/",
			f"/cars/protocols/{self.protocol.id}/measurement/",
			f"/cars/protocols/{self.protocol.id}/photos/",
			f"/cars/protocols/{self.protocol.id}/generate-docx/",
			f"/cars/protocols/{self.protocol.id}/preview-pdf/",
		]

		for url in protected_get_urls:
			with self.subTest(url=url):
				response = self.client.get(url)
				self.assertEqual(response.status_code, 403)
				self.assertEqual(response.data["code"], "permission_denied")
				self.assertIn("detail", response.data)

		self.assertEqual(
			self.client.patch(
				f"/cars/protocols/{self.protocol.id}/update/",
				{"owner_name": "Подмена"},
				format="json",
			).status_code,
			403,
		)
		self.assertEqual(
			self.client.get(f"/cars/protocol-photos/{photo.id}/file/").status_code,
			403,
		)
		self.assertEqual(
			self.client.delete(f"/cars/protocol-photos/{photo.id}/delete/").status_code,
			403,
		)

	def test_protocol_serializer_rejects_manual_service_field_changes(self):
		self.authenticate(self.measurer)
		self.assertEqual(
			self.post(f"/cars/protocols/{self.protocol.id}/start-editing/").status_code,
			200,
		)

		response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{
				"protocol_number": "HACKED-NUMBER",
				"status": "approved",
				"user": self.other_measurer.id,
				"locked_by": self.other_measurer.id,
				"locked_at": "2030-01-01T00:00:00Z",
				"cancelled_by": self.other_measurer.id,
				"owner_name": "Разрешенное изменение",
			},
			format="json",
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.protocol_number, "00001")
		self.assertEqual(self.protocol.status, "measurement")
		self.assertEqual(self.protocol.user_id, self.measurer.id)
		self.assertEqual(self.protocol.locked_by_id, self.measurer.id)
		self.assertEqual(self.protocol.owner_name, "Разрешенное изменение")


def write_fake_docx(template_path, output_path, context):
	Path(output_path).write_bytes(b"new")


class ProtocolTemplateGenerationTests(SimpleTestCase):
	def test_generator_uses_the_single_protocol_template(self):
		protocol = SimpleNamespace(id=42)

		with TemporaryDirectory() as media_root:
			with override_settings(MEDIA_ROOT=media_root):
				with patch(
					"cars.services.protocol_docx_generator.protocol_document_path",
					return_value=Path(media_root) / "00042" / "00042 Протокол.docx",
				), patch(
					"cars.services.protocol_docx_generator.build_protocol_docx_context",
					return_value={"example": "value"},
				), patch(
					"cars.services.protocol_docx_generator.render_protocol_docx",
					side_effect=write_fake_docx,
				) as render, patch(
					"cars.services.protocol_docx_generator.add_protocol_placeholder_aliases"
				) as add_aliases, patch(
					"cars.services.protocol_docx_generator.count_word_pages",
					side_effect=[40, 40],
				) as count_pages:
					output = generate_protocol_docx(protocol)

		self.assertEqual(output.name, "00042 Протокол.docx")
		self.assertEqual(add_aliases.call_count, 1)
		self.assertEqual(add_aliases.call_args.args[0]["example"], "value")
		self.assertEqual(render.call_count, 2)
		self.assertEqual(count_pages.call_count, 2)
		self.assertEqual(
			render.call_args.kwargs["template_path"].name,
			"protocol_template.docx",
		)
		self.assertEqual(
			render.call_args.kwargs["context"],
			{"example": "value", "protocol_page_count": "40"},
		)

	def render_template_with_photos(self, temp_dir, count):
		photos = []
		for index in range(count):
			image_path = Path(temp_dir) / f"{index}.png"
			Image.new("RGB", (40, 30), "red").save(image_path)
			photos.append({"path": str(image_path), "title": f"Фото слота {index}"})

		output_path = Path(temp_dir) / "output.docx"
		template_path = Path(__file__).parent / "templates" / "protocol_template.docx"
		render_protocol_docx(template_path, output_path, {"protocol_photos": photos})
		return Document(output_path)

	def test_photo_table_has_two_photos_per_row_with_captions(self):
		with TemporaryDirectory() as temp_dir:
			document = self.render_template_with_photos(temp_dir, 3)

		photo_tables = [
			table for table in document.tables
			if any("Фото слота" in cell.text for row in table.rows for cell in row.cells)
		]
		self.assertEqual(len(photo_tables), 1)
		rows = photo_tables[0].rows
		self.assertEqual(len(rows), 4)
		self.assertEqual(
			[[cell.text for cell in row.cells] for row in rows[1::2]],
			[["Фото слота 0", "Фото слота 1"], ["Фото слота 2", ""]],
		)
		self.assertEqual(
			[len(cell._tc.xpath(".//w:drawing")) for row in rows[0::2] for cell in row.cells],
			[1, 1, 1, 0],
		)
		self.assertNotIn("photo_image", document.element.xml)
		self.assertNotIn("photo_caption", document.element.xml)

	def test_photo_table_is_replaced_with_text_when_there_are_no_photos(self):
		with TemporaryDirectory() as temp_dir:
			document = self.render_template_with_photos(temp_dir, 0)

		paragraphs = [paragraph.text for paragraph in document.paragraphs]
		photos_heading = paragraphs.index("5. Фотоматериалы")
		self.assertIn("Отсутствуют.", paragraphs[photos_heading:photos_heading + 4])
		self.assertNotIn("photo_image", document.element.xml)

	def test_protocol_page_count_placeholder_is_replaced_with_word_page_count(self):
		with TemporaryDirectory() as temp_dir:
			template_path = Path(temp_dir) / "template.docx"
			output_path = Path(temp_dir) / "output.docx"
			template = Document()
			paragraph = template.add_paragraph()
			paragraph.add_run("Оформлен на ")
			paragraph.add_run("{{ protocol_page_count }}")
			paragraph.add_run(" л.")
			template.save(template_path)

			render_protocol_docx(template_path, output_path, {"protocol_page_count": "40"})
			generated = Document(output_path)

		self.assertEqual(generated.paragraphs[0].text, "Оформлен на 40 л.")
		self.assertFalse(generated.paragraphs[0]._p.xpath(".//w:instrText"))

	def test_header_document_selections_compose_one_placeholder_value(self):
		protocol = SimpleNamespace(
			header_document_selections=["application", "kuts", "sampling_data_disclaimer"]
		)

		self.assertEqual(
			build_submitted_documents(protocol),
			"Заявка, КУТС, ИЛ не несет ответственность за данные отбора образца.",
		)

	def test_application_documents_list_never_contains_application(self):
		protocol = SimpleNamespace(header_document_selections=["kuts", "application", "tpo"])

		self.assertEqual(build_submitted_documents(protocol), "Заявка, ТПО, КУТС")
		self.assertEqual(build_submitted_documents(protocol, exclude={"application"}), "ТПО, КУТС")

		only_application = SimpleNamespace(header_document_selections=["application"])
		self.assertEqual(build_submitted_documents(only_application, exclude={"application"}), "")

	def test_other_documents_text_goes_last_into_documents_list(self):
		protocol = SimpleNamespace(
			header_document_selections=["other", "application", "tpo"],
			header_documents_other="  доверенность\nот 01.10.2026  ",
		)

		self.assertEqual(build_submitted_documents(protocol), "Заявка, ТПО, доверенность\nот 01.10.2026")
		self.assertEqual(
			build_submitted_documents(protocol, exclude={"application"}),
			"ТПО, доверенность\nот 01.10.2026",
		)

		empty_other = SimpleNamespace(header_document_selections=["tpo", "other"], header_documents_other="  ")
		self.assertEqual(build_submitted_documents(empty_other), "ТПО")

		unchecked_other = SimpleNamespace(header_document_selections=["tpo"], header_documents_other="доверенность")
		self.assertEqual(build_submitted_documents(unchecked_other), "ТПО")

	def test_protocol_context_aliases_supply_template_fields(self):
		placeholder_aliases, condition_aliases = load_placeholder_aliases()

		context = {name: name for name in placeholder_aliases}
		for source_tag in condition_aliases:
			match = re.fullmatch(
				r"{%\s*tr\s+if\s+(?:not\s+)?([A-Za-z0-9_]+)\s*%}",
				source_tag,
			)
			if match:
				context.setdefault(match.group(1), False)
		context["rear_fog_a_8_13_2_conclusion"] = "-"
		context["applicable_6_4"] = (
			"не применяется (сервоприводы для включения и/или отключения устройства "
			"для предотвращения несанкционированного использования не используются)"
		)

		add_protocol_placeholder_aliases(context)
		self.assertEqual(context["a_8_13_2_conclusion"], "-")
		# А.6.4 не применяется: причина в статусе, в заключении прочерк.
		self.assertEqual(context["a_6_4_conclusion"], "-")
		for canonical_name in placeholder_aliases.values():
			self.assertIn(canonical_name, context)

	def test_conclusions_are_dashed_for_not_applicable_and_electric_clauses(self):
		electric_codes = (
			"a_20_1", "a_20_2", "a_20_3", "a_20_5", "a_20_5_3", "a_20_6", "a_20_7", "a_20_8",
			"a_21_4", "a_21_5", "a_21_6", "a_22_3", "a_22_4",
		)
		context = {
			"applicable_18_4_2": "не применяется (на ТС отсутствуют выступающие элементы крыши из жесткого материала)",
			"applicable_18_4_4": "соответствует",
		}

		electric = build_dash_conclusion_values(context, SimpleNamespace(fuel_type="electric"))
		petrol = build_dash_conclusion_values(context, SimpleNamespace(fuel_type="petrol"))

		self.assertEqual(electric["a_18_4_2_conclusion"], "-")
		self.assertEqual(
			petrol["a_18_4_4_conclusion"],
			"Соответствует требованиям\nТР ТС 018/2011\nПриложение №4 п.3.4.4.4",
		)
		for code in electric_codes:
			with self.subTest(code=code):
				self.assertEqual(electric[f"{code}_conclusion"], "-")
				self.assertTrue(petrol[f"{code}_conclusion"].startswith("Соответствует требованиям"))
		self.assertEqual(
			petrol["a_22_4_conclusion"],
			"Соответствует требованиям\nТР ТС 018/2011\nПриложение №8 п.9.7",
		)

	def test_co_uncertainty_comes_from_gas_analyzer_error(self):
		values = build_eco_values(
			SimpleNamespace(),
			SimpleNamespace(fuel_type="petrol", mileage_km=5000, co_min_pct=Decimal("0.10"), co_max_pct=Decimal("0.05")),
		)
		# Газоанализатор ±0,03 %: U = 1,65 · 0,03 / √3 ≈ 0,03 %.
		self.assertEqual(values["u_co_min_21_7"], "0,03")
		self.assertEqual(values["co_min_result_21_7"], "0,10 % ± 0,03 %")
		self.assertIn("Минимальная - 0,10 % ± 0,03 %", values["full_result_a_21_7"])
		self.assertIn("Повышенная - 0,05 % ± 0,03 %", values["full_result_a_21_7"])
		self.assertNotIn("0,19", values["full_result_a_21_7"])

	def test_electric_vehicle_has_dash_in_21_9_conclusion(self):
		values = build_eco_values(SimpleNamespace(), SimpleNamespace(fuel_type="electric", mileage_km=None))
		self.assertEqual(values["full_result_a_21_9"], "-")

	def test_template_has_placeholders_for_dashed_conclusions(self):
		template = Path(__file__).parent / "templates" / "protocol_template.docx"
		xml = re.sub(r"<[^>]+>", "", zipfile.ZipFile(template).read("word/document.xml").decode("utf-8"))
		for code in STATIC_CONCLUSIONS:
			with self.subTest(code=code):
				self.assertIn(f"{{{{ {code}_conclusion }}}}", xml)

	def test_missing_registration_number_marks_clause_1_2_1_not_applicable(self):
		protocol = SimpleNamespace(registration_number=None)
		values = build_applicability_values(
			protocol,
			SimpleNamespace(),
			SimpleNamespace(),
			{},
		)

		self.assertEqual(
			values["applicable_1_2_1"],
			"не применяется (на ТС отсутствует государственный регистрационный знак)",
		)


class ProtocolFileStorageTests(TestCase):
	def setUp(self):
		temp_dir = TemporaryDirectory()
		self.addCleanup(temp_dir.cleanup)
		self.temp = Path(temp_dir.name)
		self.media_root = self.temp / "media"
		self.in_progress = self.temp / "in_progress"
		self.completed = self.temp / "completed"

		media_override = override_settings(MEDIA_ROOT=str(self.media_root))
		media_override.enable()
		self.addCleanup(media_override.disable)

		ProtocolStorageSettings.objects.create(
			in_progress_dir=str(self.in_progress),
			completed_dir=str(self.completed),
		)

		self.client = APIClient()
		self.measurer = User.objects.create_user(username="measurer", password="test-password")
		self.measurer.groups.add(Group.objects.get_or_create(name="measurer")[0])
		self.manager = User.objects.create_user(username="manager", password="test-password")
		self.manager.groups.add(Group.objects.get_or_create(name="manager")[0])
		self.superuser = User.objects.create_superuser(username="root", password="test-password")
		self.protocol = Protocol.objects.create(
			user=self.measurer,
			protocol_number="00001",
			vin="XW8ZZZ61ZEG000001",
			status="measurement",
			locked_by=self.measurer,
			locked_at=timezone.now(),
		)
		# Папка «<номер> <VIN> <дата создания>», как при создании через сайт.
		self.folder = f"00001 XW8ZZZ61ZEG000001 {timezone.now().astimezone().strftime('%d.%m.%Y')}"
		self.protocol.files_folder = self.folder
		self.protocol.save(update_fields=["files_folder"])

	def post(self, path, data=None):
		return self.client.post(path, data or {}, format="json")

	def make_protocol_file(self, root, name="photos/old.jpg", content=b"old"):
		path = root / self.folder / name
		path.parent.mkdir(parents=True, exist_ok=True)
		path.write_bytes(content)
		return path

	def upload_photo(self, photo_type="odometer_photo", filename="camera.JPG", content=b"jpeg-bytes"):
		self.client.force_authenticate(user=self.measurer)
		return self.client.post(
			f"/cars/protocols/{self.protocol.id}/photos/create/",
			{
				"photo_type": photo_type,
				"file": SimpleUploadedFile(filename, content, content_type="image/jpeg"),
			},
			format="multipart",
		)

	def test_created_protocol_gets_folder_named_by_number_vin_and_date(self):
		self.client.force_authenticate(user=self.measurer)
		response = self.post("/cars/protocols/create/", {"owner_name": "Владелец", "vin": "JTDBR32E720123456"})

		self.assertEqual(response.status_code, 201)
		today = timezone.now().astimezone().strftime("%d.%m.%Y")
		folder = f"{response.data['protocol_number']} JTDBR32E720123456 {today}"
		self.assertEqual(response.data["files_folder"], folder)
		self.assertTrue((self.in_progress / folder).is_dir())

	def test_vin_change_renames_protocol_folder(self):
		self.upload_photo()
		self.client.force_authenticate(user=self.measurer)
		response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"vin": "NEWVIN0000000001"},
			format="json",
		)

		self.assertEqual(response.status_code, 200)
		self.assertIsNone(response.data["files_warning"])
		new_folder = self.folder.replace("XW8ZZZ61ZEG000001", "NEWVIN0000000001")
		self.assertEqual(response.data["files_folder"], new_folder)
		self.assertFalse((self.in_progress / self.folder).exists())
		self.assertTrue((self.in_progress / new_folder / "photos" / "Одометр.jpg").exists())

		photo = ProtocolPhoto.objects.get(protocol=self.protocol)
		file_response = self.client.get(f"/cars/protocol-photos/{photo.id}/file/")
		self.assertEqual(file_response.status_code, 200)
		file_response.close()

	def test_vin_change_keeps_old_folder_when_target_exists(self):
		self.make_protocol_file(self.in_progress)
		new_folder = self.folder.replace("XW8ZZZ61ZEG000001", "NEWVIN0000000001")
		(self.in_progress / new_folder).mkdir(parents=True)

		self.client.force_authenticate(user=self.measurer)
		response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"vin": "NEWVIN0000000001"},
			format="json",
		)

		self.assertEqual(response.status_code, 200)
		self.assertIn("уже существует", response.data["files_warning"])
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.vin, "NEWVIN0000000001")
		self.assertEqual(self.protocol.files_folder, self.folder)
		self.assertTrue((self.in_progress / self.folder / "photos" / "old.jpg").exists())

	def test_files_folder_cannot_be_changed_through_api(self):
		self.client.force_authenticate(user=self.measurer)
		self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"files_folder": "..\hack", "files_location": "completed"},
			format="json",
		)

		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.files_folder, self.folder)
		self.assertEqual(self.protocol.files_location, "in_progress")

	def test_relocate_command_renames_old_number_folders_by_vin(self):
		Protocol.objects.filter(pk=self.protocol.pk).update(files_folder="00001")
		old = self.in_progress / "00001" / "photos" / "old.jpg"
		old.parent.mkdir(parents=True)
		old.write_bytes(b"old")

		call_command("relocate_protocol_files", stdout=StringIO())

		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.files_folder, self.folder)
		self.assertTrue((self.in_progress / self.folder / "photos" / "old.jpg").exists())
		self.assertFalse((self.in_progress / "00001").exists())

	def test_new_protocol_gets_50_hz_network_frequency_that_can_be_changed(self):
		self.client.force_authenticate(user=self.measurer)
		response = self.post("/cars/protocols/create/", {"vin": "JTDBR32E720123456"})
		self.assertEqual(response.status_code, 201)

		protocol = Protocol.objects.get(pk=response.data["id"])
		self.assertEqual(protocol.power_supply.frequency_hz, Decimal("50"))

		protocol.locked_by = self.measurer
		protocol.locked_at = timezone.now()
		protocol.save(update_fields=["locked_by", "locked_at"])
		update = self.client.patch(
			f"/cars/protocols/{protocol.id}/power-supply/update/",
			{"frequency_hz": "49.5"},
			format="json",
		)
		self.assertEqual(update.status_code, 200, update.data)
		protocol.power_supply.refresh_from_db()
		self.assertEqual(protocol.power_supply.frequency_hz, Decimal("49.5"))

	def test_protocol_created_without_owner_name_keeps_it_empty(self):
		self.client.force_authenticate(user=self.measurer)
		response = self.post("/cars/protocols/create/", {"brand_name": "KIA"})

		self.assertEqual(response.status_code, 201)
		self.assertEqual(response.data["owner_name"], "")

	def test_owner_legal_entity_is_saved_and_goes_to_docx_context(self):
		legal_entity = "ООО «Ромашка»\nг. Владивосток, ул. Снеговая, 64"
		self.client.force_authenticate(user=self.measurer)
		response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"owner_legal_entity": legal_entity},
			format="json",
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.owner_legal_entity, legal_entity)
		self.assertEqual(build_protocol_docx_context(self.protocol)["owner_legal_entity"], legal_entity)

	def test_owner_requisites_string_goes_to_docx_context_as_is(self):
		requisites = "ОГРН 1234567890123, ИНН/КПП 2540000000/254001001,\nр/с 40702810000000000000"
		self.client.force_authenticate(user=self.measurer)
		response = self.client.patch(
			f"/cars/protocols/{self.protocol.id}/update/",
			{"owner_requisites": requisites},
			format="json",
		)

		self.assertEqual(response.status_code, 200)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.owner_requisites, requisites)
		self.assertEqual(build_protocol_docx_context(self.protocol)["owner_requisites"], requisites)

	def test_uploaded_photo_goes_to_in_progress_folder_and_is_served(self):
		response = self.upload_photo()

		self.assertEqual(response.status_code, 201)
		file_path = response.data["file_path"]
		self.assertEqual(file_path, "photos/Одометр.jpg")
		self.assertEqual(response.data["caption"], "Одометр")
		self.assertEqual((self.in_progress / self.folder / file_path).read_bytes(), b"jpeg-bytes")
		self.assertEqual(
			response.data["file_url"],
			f"/cars/protocol-photos/{response.data['id']}/file/",
		)

		file_response = self.client.get(response.data["file_url"])
		self.assertEqual(file_response.status_code, 200)
		self.assertEqual(file_response["Content-Type"], "image/jpeg")
		self.assertEqual(b"".join(file_response.streaming_content), b"jpeg-bytes")
		file_response.close()

	def test_photo_file_is_protected_from_foreign_measurer(self):
		photo_id = self.upload_photo().data["id"]
		other = User.objects.create_user(username="other", password="test-password")
		other.groups.add(Group.objects.get_or_create(name="measurer")[0])

		self.client.force_authenticate(user=other)
		self.assertEqual(self.client.get(f"/cars/protocol-photos/{photo_id}/file/").status_code, 403)

		self.client.force_authenticate(user=None)
		self.assertEqual(self.client.get(f"/cars/protocol-photos/{photo_id}/file/").status_code, 401)

	def test_new_photo_replaces_slot_photo_and_its_file(self):
		first = self.upload_photo("stand_test_photo", "a.png").data
		second = self.upload_photo("stand_test_photo", "b.jpg", b"second").data
		third = self.upload_photo("stand_test_photo", "c.jpg", b"third").data

		self.assertEqual(first["file_path"], "photos/Испытания на тормозном стенде.png")
		self.assertFalse((self.in_progress / self.folder / first["file_path"]).exists())
		self.assertEqual(second["file_path"], third["file_path"])
		self.assertEqual((self.in_progress / self.folder / third["file_path"]).read_bytes(), b"third")
		self.assertEqual(ProtocolPhoto.objects.filter(protocol=self.protocol).count(), 1)

	def test_only_photo_slots_are_accepted_and_listed(self):
		self.assertEqual(self.upload_photo("other").status_code, 400)
		self.assertEqual(self.upload_photo("").status_code, 400)

		ProtocolPhoto.objects.create(protocol=self.protocol, photo_type="other", file_path="x.jpg")
		self.upload_photo("speed_photo")
		self.upload_photo("acceptance_photo")

		listed = self.client.get(f"/cars/protocols/{self.protocol.id}/photos/").data
		self.assertEqual([photo["photo_type"] for photo in listed], ["acceptance_photo", "speed_photo"])

		full = self.client.get(f"/cars/protocols/{self.protocol.id}/full/").data
		self.assertEqual([photo["photo_type"] for photo in full["photos"]], ["acceptance_photo", "speed_photo"])

	def test_form_settings_require_all_fields_by_default_and_follow_admin(self):
		self.client.force_authenticate(user=self.measurer)
		self.assertEqual(self.client.get("/cars/protocol-form-settings/").data, {"require_all_fields": True})

		ProtocolFormSettings.objects.create(require_all_fields=False)
		self.assertEqual(self.client.get("/cars/protocol-form-settings/").data, {"require_all_fields": False})

		self.client.force_authenticate(user=None)
		self.assertEqual(self.client.get("/cars/protocol-form-settings/").status_code, 401)

	def test_photo_slots_follow_excel_scheme(self):
		self.client.force_authenticate(user=self.measurer)
		slots = self.client.get("/cars/protocol-photo-slots/").data

		self.assertEqual(len(slots), 31)
		self.assertEqual(slots[0], {
			"photo_type": "acceptance_photo",
			"title": "Приемка ТС",
			"description": "Эксперт с открытым капотом ТС",
			"group": "after_conditions",
		})
		groups = [slot["group"] for slot in slots]
		self.assertEqual(
			[groups.count(group) for group in (
				"after_conditions",
				"after_steering_transmission",
				"after_brakes",
				"after_lights_geometry",
			)],
			[11, 1, 13, 6],
		)
		self.assertEqual(slots[-1]["title"], "Измерение скорости ТС")
		self.assertEqual(len({slot["photo_type"] for slot in slots}), 31)

	def test_photo_file_name_is_safe_for_windows(self):
		self.assertEqual(photo_file_stem("Проверка: VIN/номер?"), "Проверка_ VIN_номер_")
		self.assertEqual(photo_file_stem(" . "), "Фото")

	def test_approve_moves_folder_to_completed_and_return_moves_it_back(self):
		photo_id = self.upload_photo().data["id"]
		Protocol.objects.filter(pk=self.protocol.pk).update(status="review", locked_by=None)

		self.client.force_authenticate(user=self.manager)
		self.assertEqual(self.post(f"/cars/protocols/{self.protocol.id}/approve/").status_code, 200)

		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "approved")
		self.assertEqual(self.protocol.files_location, "completed")
		self.assertFalse((self.in_progress / self.folder).exists())
		self.assertTrue((self.completed / self.folder / "photos").is_dir())
		file_response = self.client.get(f"/cars/protocol-photos/{photo_id}/file/")
		self.assertEqual(file_response.status_code, 200)
		# Открытый файл на Windows не даст перенести папку.
		file_response.close()

		self.assertEqual(self.post(f"/cars/protocols/{self.protocol.id}/return-to-draft/").status_code, 200)

		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "operator")
		self.assertEqual(self.protocol.files_location, "in_progress")
		self.assertFalse((self.completed / self.folder).exists())
		self.assertTrue((self.in_progress / self.folder / "photos").is_dir())

	def test_approve_returns_folder_back_when_status_is_not_saved(self):
		self.make_protocol_file(self.in_progress)
		Protocol.objects.filter(pk=self.protocol.pk).update(status="review", locked_by=None)

		self.client.force_authenticate(user=self.manager)
		with patch.object(Protocol, "save", side_effect=RuntimeError("db is down")):
			response = self.post(f"/cars/protocols/{self.protocol.id}/approve/")

		self.assertEqual(response.status_code, 500)
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "review")
		self.assertEqual(self.protocol.files_location, "in_progress")
		self.assertTrue((self.in_progress / self.folder / "photos" / "old.jpg").exists())
		self.assertFalse((self.completed / self.folder).exists())

	def test_approve_fails_without_merge_when_completed_folder_exists(self):
		self.make_protocol_file(self.in_progress)
		self.make_protocol_file(self.completed, name="other.txt")
		Protocol.objects.filter(pk=self.protocol.pk).update(status="review", locked_by=None)

		self.client.force_authenticate(user=self.manager)
		response = self.post(f"/cars/protocols/{self.protocol.id}/approve/")

		self.assertEqual(response.status_code, 409)
		self.assertIn("уже существует", response.data["detail"])
		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.status, "review")
		self.assertTrue((self.in_progress / self.folder / "photos" / "old.jpg").exists())
		self.assertFalse((self.completed / self.folder / "photos").exists())

	def test_cancelled_protocol_folder_stays_in_progress(self):
		self.make_protocol_file(self.in_progress)
		Protocol.objects.filter(pk=self.protocol.pk).update(status="review", locked_by=None)

		self.client.force_authenticate(user=self.manager)
		self.assertEqual(self.post(f"/cars/protocols/{self.protocol.id}/cancel/").status_code, 200)

		self.protocol.refresh_from_db()
		self.assertEqual(self.protocol.files_location, "in_progress")
		self.assertTrue((self.in_progress / self.folder / "photos" / "old.jpg").exists())

	def test_delete_protocol_removes_its_folder(self):
		self.make_protocol_file(self.in_progress)

		self.client.force_authenticate(user=self.superuser)
		response = self.client.delete(f"/cars/protocols/{self.protocol.id}/delete/")

		self.assertEqual(response.status_code, 204)
		self.assertFalse((self.in_progress / self.folder).exists())
		self.assertTrue(self.in_progress.is_dir())

	def test_admin_clear_removes_only_protocol_folders(self):
		from .admin import clear_protocol_data

		self.make_protocol_file(self.in_progress)
		unrelated = self.in_progress / "чужая папка" / "file.txt"
		unrelated.parent.mkdir(parents=True)
		unrelated.write_bytes(b"keep")

		clear_protocol_data()

		self.assertFalse((self.in_progress / self.folder).exists())
		self.assertTrue(unrelated.exists())

	def make_document_templates(
		self, templates_dir, titles=("protocol", "application", "tech_record"), evaluation_sheet=True,
	):
		for name in titles:
			template = Document()
			template.add_paragraph(f"{name}: {{{{ protocol_number }}}} {{{{ applicant_name }}}}")
			template.save(templates_dir / f"{name}_template.docx")
		if evaluation_sheet:
			shutil.copy(
				Path(__file__).parent / "templates" / "evaluation_sheet_template.xlsx",
				templates_dir / "evaluation_sheet_template.xlsx",
			)

	def generate_documents(self, templates_dir):
		self.client.force_authenticate(user=self.measurer)
		with patch("cars.services.protocol_docx_generator.TEMPLATES_DIR", templates_dir), patch(
			"cars.services.protocol_docx_generator.count_word_pages",
			return_value=1,
		) as count_pages:
			response = self.post(f"/cars/protocols/{self.protocol.id}/generate-docx/")
		return response, count_pages

	def test_generate_builds_all_documents_into_protocol_folder(self):
		templates_dir = self.temp / "templates"
		templates_dir.mkdir()
		self.make_document_templates(templates_dir)
		legacy_docx = self.in_progress / self.folder / "00001.docx"
		legacy_docx.parent.mkdir(parents=True)
		legacy_docx.write_bytes(b"old")
		Protocol.objects.filter(pk=self.protocol.pk).update(
			owner_name="Иванов И. И.",
			owner_legal_entity="ООО «Ромашка»",
		)

		response, count_pages = self.generate_documents(templates_dir)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.data["folder"], str(self.in_progress / self.folder))
		self.assertEqual(
			response.data["files"],
			["00001 Протокол.docx", "00001 Заявка.docx", "00001 Техзапись.docx", "00001 Лист оценки.xlsx"],
		)
		self.assertEqual(response.data["skipped"], [])
		# Страницы через Word считаются только для протокола.
		self.assertEqual(count_pages.call_count, 2)
		self.assertEqual(
			Document(self.in_progress / self.folder / "00001 Заявка.docx").paragraphs[0].text,
			"application: 00001 ООО «Ромашка»",
		)
		self.assertEqual(
			sorted(path.name for path in (self.in_progress / self.folder).iterdir()),
			["00001 Заявка.docx", "00001 Лист оценки.xlsx", "00001 Протокол.docx", "00001 Техзапись.docx"],
		)
		sheet = load_workbook(self.in_progress / self.folder / "00001 Лист оценки.xlsx").active
		self.assertTrue(sheet["B1"].value.startswith("Лист оценки неопределенности к протоколу № 00001"))

	def test_generate_skips_documents_without_template(self):
		templates_dir = self.temp / "templates"
		templates_dir.mkdir()
		self.make_document_templates(templates_dir, titles=("protocol",), evaluation_sheet=False)

		response, _ = self.generate_documents(templates_dir)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.data["files"], ["00001 Протокол.docx"])
		self.assertEqual(
			response.data["skipped"],
			[
				"Заявка (нет шаблона application_template.docx)",
				"Техзапись (нет шаблона tech_record_template.docx)",
				"Лист оценки (нет шаблона evaluation_sheet_template.xlsx)",
			],
		)

	def test_busy_document_is_reported_and_others_are_still_built(self):
		templates_dir = self.temp / "templates"
		templates_dir.mkdir()
		self.make_document_templates(templates_dir)
		folder = self.in_progress / self.folder
		folder.mkdir(parents=True)
		busy = folder / "00001 Заявка.docx"
		busy.write_bytes(b"opened in word")
		real_replace = os.replace

		def replace(source, target):
			if Path(target) == busy:
				raise PermissionError("locked")
			real_replace(source, target)

		with patch("cars.services.protocol_storage.os.replace", side_effect=replace):
			response, _ = self.generate_documents(templates_dir)

		self.assertEqual(response.status_code, 409)
		self.assertIn("занят", response.data["detail"])
		self.assertEqual(
			response.data["files"],
			["00001 Протокол.docx", "00001 Техзапись.docx", "00001 Лист оценки.xlsx"],
		)
		self.assertEqual(busy.read_bytes(), b"opened in word")
		self.assertEqual(
			sorted(path.name for path in folder.iterdir()),
			["00001 Заявка.docx", "00001 Лист оценки.xlsx", "00001 Протокол.docx", "00001 Техзапись.docx"],
		)

	def test_manufacture_year_comes_from_selected_manufacture_date(self):
		self.protocol.manufacture_year = 2013
		self.assertEqual(build_protocol_docx_context(self.protocol)["manufacture_year"], "2013")

		self.protocol.manufacture_date = date(2014, 11, 1)
		self.assertEqual(build_protocol_docx_context(self.protocol)["manufacture_year"], "2014")

	def test_applicant_name_falls_back_to_owner_full_name(self):
		self.protocol.owner_name = "Иванов И. И."
		self.assertEqual(build_protocol_docx_context(self.protocol)["applicant_name"], "Иванов И. И.")

		self.protocol.owner_legal_entity = "ООО «Ромашка»"
		self.assertEqual(build_protocol_docx_context(self.protocol)["applicant_name"], "ООО «Ромашка»")

	def test_folder_name_falls_back_to_padded_id(self):
		self.assertEqual(protocol_number_label(SimpleNamespace(id=7, protocol_number="A-15")), "A-15")
		for number in (None, "", "TEMP-abc", "12/34", "bad:name", "dot."):
			with self.subTest(number=number):
				self.assertEqual(
					protocol_number_label(SimpleNamespace(id=7, protocol_number=number)),
					"00007",
				)

	def test_storage_dirs_validation(self):
		validate_storage_dirs(str(self.temp / "a"), str(self.temp / "b"))
		validate_storage_dirs("", "")
		self.assertTrue((self.media_root / "protocols" / "completed").is_dir())

		invalid = [
			("relative/path", str(self.temp / "b")),
			(str(self.temp / "same"), str(self.temp / "same")),
			(str(self.temp / "outer"), str(self.temp / "outer" / "inner")),
		]
		for in_progress, completed in invalid:
			with self.subTest(in_progress=in_progress, completed=completed):
				with self.assertRaises(ValidationError):
					validate_storage_dirs(in_progress, completed)

	def test_relocate_command_moves_legacy_photos(self):
		legacy = self.media_root / "protocol_photos" / str(self.protocol.id) / "other_1.jpg"
		legacy.parent.mkdir(parents=True)
		legacy.write_bytes(b"legacy")
		moved = ProtocolPhoto.objects.create(
			protocol=self.protocol,
			file_path=f"protocol_photos/{self.protocol.id}/other_1.jpg",
		)
		missing = ProtocolPhoto.objects.create(
			protocol=self.protocol,
			file_path=f"protocol_photos/{self.protocol.id}/lost.jpg",
		)

		call_command("relocate_protocol_files", "--dry-run", stdout=StringIO())
		moved.refresh_from_db()
		self.assertTrue(legacy.exists())
		self.assertTrue(moved.file_path.startswith("protocol_photos/"))

		output = StringIO()
		call_command("relocate_protocol_files", stdout=output)
		moved.refresh_from_db()
		self.assertEqual(moved.file_path, "photos/other_1.jpg")
		self.assertEqual((self.in_progress / self.folder / moved.file_path).read_bytes(), b"legacy")
		self.assertFalse(legacy.exists())
		self.assertIn(f"фото id={missing.id}", output.getvalue())

	def test_relocate_command_creates_folders_for_protocols_without_files(self):
		self.assertFalse((self.in_progress / self.folder).exists())

		call_command("relocate_protocol_files", "--dry-run", stdout=StringIO())
		self.assertFalse((self.in_progress / self.folder).exists())

		call_command("relocate_protocol_files", stdout=StringIO())
		self.assertTrue((self.in_progress / self.folder).is_dir())

	def test_relocate_command_moves_folders_from_previous_root(self):
		old_root = self.temp / "old_in_progress"
		self.make_protocol_file(old_root)
		photo = ProtocolPhoto.objects.create(protocol=self.protocol, file_path="photos/old.jpg")

		call_command("relocate_protocol_files", "--source", str(old_root), stdout=StringIO())

		self.assertFalse((old_root / self.folder).exists())
		self.assertTrue((self.in_progress / self.folder / photo.file_path).exists())


class EvaluationSheetTests(SimpleTestCase):
	def make_template(self, path, rows):
		workbook = Workbook()
		sheet = workbook.active
		for index, row in enumerate(rows, start=1):
			for column, value in enumerate(row, start=1):
				sheet.cell(row=index, column=column, value=value)
		sheet["B2"].number_format = "0.00"
		workbook.save(path)

	def test_xlsx_renderer_fills_numbers_and_removes_rows_by_conditions(self):
		with TemporaryDirectory() as temp_dir:
			template = Path(temp_dir) / "template.xlsx"
			output = Path(temp_dir) / "output.xlsx"
			self.make_template(template, [
				["Протокол № {{ number }}", None],
				["число", "{{ value }}"],
				["{%tr if one %}", None],
				["один", None],
				["{%tr elif two %}", None],
				["два", None],
				["{%tr else %}", None],
				["нет", None],
				["{%tr endif %}", None],
				["после", "{{ missing }}"],
			])

			render_xlsx(template, output, {"number": "00012", "value": 490.0, "one": False, "two": True})
			sheet = load_workbook(output).active

		self.assertEqual(
			[sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)],
			["Протокол № 00012", "число", "два", "после"],
		)
		self.assertEqual(sheet["B2"].value, 490)
		self.assertEqual(sheet["B2"].number_format, "0.00")
		self.assertIn(sheet["B4"].value, ("", None))

	def make_protocol(self, fuel_type="petrol", rear_fog_count=1):
		return SimpleNamespace(
			protocol_number="00012",
			protocol_date=date(2026, 10, 29),
			tire_season="summer",
			measurement=SimpleNamespace(
				fuel_type=fuel_type,
				vehicle_length_mm=Decimal("4640"),
				vehicle_width_mm=Decimal("1780"),
				vehicle_height_mm=Decimal("1435"),
				co_min_pct=Decimal("0.1"),
				co_max_pct=None,
				sun_strip_width_mm=None,
				light_absorption_1=Decimal("0.20"),
				light_absorption_2=Decimal("0.22"),
				light_absorption_3=Decimal("0.24"),
				light_absorption_4=None,
				light_absorption_5=None,
				light_absorption_6=None,
				axle1_load_kg=Decimal("995"),
			),
			brake=SimpleNamespace(
				parking_brake_type="mechanical_hand",
				service_brake_control_force_axle1_n=Decimal("490"),
			),
			light=SimpleNamespace(rear_fog_count=rear_fog_count, front_fog_count=0),
		)

	def test_evaluation_values_match_excel_prototype(self):
		context = build_evaluation_context(self.make_protocol())

		self.assertEqual(context["protocol_number"], "00012")
		# Рулетка: длина 4640 мм — интервал 4–5 м, Δ = 0,75 мм.
		self.assertEqual(context["ev_vehicle_length_error"], "±0,75 мм")
		self.assertAlmostEqual(context["ev_vehicle_length_ub"], 0.4330127, places=6)
		self.assertAlmostEqual(context["ev_vehicle_length_u"], 0.7144710, places=6)
		self.assertEqual(context["ev_vehicle_width_error"], "±0,3 мм")
		self.assertAlmostEqual(context["ev_vehicle_width_u"], 0.2857884, places=6)
		self.assertAlmostEqual(context["ev_vehicle_height_ub"], 1.1902381, places=6)
		# Стенд: 5 % от 490 Н плюс цена деления 1 Н.
		self.assertAlmostEqual(context["ev_service_control_force_axle1_ub"], 14.1480270, places=6)
		# Не измерено — прочерк.
		self.assertEqual(context["ev_co_max_x"], "-")
		self.assertEqual(context["ev_co_max_u"], "-")
		# Дымность: 3 замера, uA по типу А.
		self.assertEqual(context["ev_light_absorption_n"], 3)
		self.assertAlmostEqual(context["ev_light_absorption_x"], 0.22, places=6)
		self.assertAlmostEqual(context["ev_light_absorption_ua"], 0.0115470, places=6)
		self.assertEqual(context["ev_vehicle_length_33987_u"], context["ev_vehicle_length_u"])

	def test_evaluation_conditions_follow_vehicle(self):
		petrol = build_evaluation_context(self.make_protocol())
		diesel = build_evaluation_context(self.make_protocol("hybrid_diesel", rear_fog_count=2))
		electric = build_evaluation_context(self.make_protocol("electric", rear_fog_count=0))

		self.assertEqual((petrol["ev_petrol_like"], petrol["ev_diesel_like"]), (True, False))
		self.assertEqual((diesel["ev_petrol_like"], diesel["ev_diesel_like"]), (False, True))
		self.assertFalse(electric["ev_not_electric"])
		self.assertEqual((petrol["ev_rear_fog_present"], petrol["ev_rear_fog_two_lamps"]), (True, False))
		self.assertTrue(diesel["ev_rear_fog_two_lamps"])
		self.assertFalse(electric["ev_rear_fog_present"])
		self.assertFalse(petrol["ev_front_fog_present"])
		self.assertFalse(petrol["ev_winter_tires"])
		self.assertTrue(petrol["ev_parking_control_force_applicable"])

		# Ширина полосы 0 — полосы нет, как в протоколе.
		for width, present in ((None, False), (Decimal("0"), False), (Decimal("90"), True)):
			with self.subTest(sun_strip=width):
				protocol = self.make_protocol()
				protocol.measurement.sun_strip_width_mm = width
				self.assertEqual(build_evaluation_context(protocol)["ev_sun_strip_present"], present)

	def test_template_has_placeholders_for_every_row(self):
		sheet = load_workbook(Path(__file__).parent / "templates" / "evaluation_sheet_template.xlsx").active
		text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)
		for row in EVALUATION_ROWS:
			with self.subTest(row=row.key):
				for suffix in ("x", "ua", "ub", "uc", "u", "error", "n"):
					self.assertIn(f"{{{{ ev_{row.key}_{suffix} }}}}", text)

class ProtocolApplicabilityTests(SimpleTestCase):
	def test_tire_depth_1072_is_always_output_and_1073_is_winter_only(self):
		measurement = SimpleNamespace(
			tire_depth_fl_mm=Decimal("5.6"),
			tire_depth_fr_mm=Decimal("5.4"),
			tire_depth_rl_mm=Decimal("5.2"),
			tire_depth_rr_mm=Decimal("5.0"),
		)
		cases = (
			("summer", True, False),
			("winter", False, True),
		)

		for season, summer_present, winter_present in cases:
			with self.subTest(season=season):
				values = build_tire_depth_values(
					SimpleNamespace(tire_season=season),
					measurement,
				)

				self.assertEqual(values["summer_tires_present"], summer_present)
				self.assertEqual(values["winter_tires_present"], winter_present)
				self.assertEqual(values["tire_depth_fl_10_7_2"], "5,6 мм")
				self.assertIn("5,6 мм ± 0,05 мм", values["full_result_a_10_7_2"])
				self.assertNotIn("мм мм", values["full_result_a_10_7_2"])
				self.assertNotIn("мм мм", values["full_result_a_10_7_3"])
				if season == "winter":
					self.assertIn("Переднее левое: 5,6 мм ± 0,05 мм", values["full_result_a_10_7_3"])

				if season == "summer":
					self.assertEqual(
						values["full_result_a_10_7_3"],
						"не применяется (на ТС установлены летние шины)",
					)
				else:
					self.assertEqual(values["tire_depth_fl_10_7_3"], "5,6 мм")
					self.assertIn("5,6 мм ± 0,05 мм", values["full_result_a_10_7_3"])

	def test_front_fog_clause_8_10_1_splits_status_and_conclusion(self):
		absent = build_front_fog_values(SimpleNamespace(front_fog_count=0))
		self.assertEqual(
			absent["front_fog_a_8_10_1_status"],
			"не применяется (в ТС отсутствуют передние противотуманные фары)",
		)
		self.assertEqual(absent["front_fog_a_8_10_1_conclusion"], "-")
		for field in (
			"fog_light_left_distance_8_10_1",
			"fog_light_right_distance_8_10_1",
			"u_fog_light_left_distance_8_10_1",
			"u_fog_light_right_distance_8_10_1",
			"fog_light_lower_point_8_10_2",
			"fog_light_upper_point_8_10_2",
			"u_fog_light_lower_point_8_10_2",
			"u_fog_light_upper_point_8_10_2",
		):
			self.assertEqual(absent[field], "")
		for clause in ("8_10_2", "8_10_3"):
			self.assertEqual(
				absent[f"front_fog_a_{clause}_status"],
				"не применяется (в ТС отсутствуют передние противотуманные фары)",
			)
			self.assertEqual(absent[f"front_fog_a_{clause}_conclusion"], "-")

		present = build_front_fog_values(SimpleNamespace(
			front_fog_count=2,
			fog_light_left_distance_mm=200,
			fog_light_right_distance_mm=210,
		))
		self.assertEqual(present["front_fog_a_8_10_1_status"], "соответствует")
		self.assertIn("Результат:", present["front_fog_a_8_10_1_conclusion"])
		self.assertIn("±", present["front_fog_a_8_10_1_conclusion"])
		self.assertIn("Левая", present["front_fog_a_8_10_1_conclusion"])
		self.assertIn("Правая", present["front_fog_a_8_10_1_conclusion"])
		self.assertEqual(present["front_fog_a_8_10_2_status"], "соответствует")
		self.assertIn("Результат:", present["front_fog_a_8_10_2_conclusion"])
		self.assertIn("±", present["front_fog_a_8_10_2_conclusion"])
		self.assertEqual(present["front_fog_a_8_10_3_status"], "соответствует")
		self.assertNotEqual(present["front_fog_a_8_10_3_conclusion"], "-")

	def test_rear_fog_condition_placeholders_cover_all_count_states(self):
		cases = (
			(None, False, False, True),
			(0, False, False, True),
			(1, True, False, False),
			(2, False, True, False),
		)

		for count, one_lamp, two_lamps, absent in cases:
			with self.subTest(rear_fog_count=count):
				values = build_light_device_row_values(
					SimpleNamespace(rear_fog_count=count)
				)

				self.assertEqual(values["rear_fog_one_lamp_present"], one_lamp)
				self.assertEqual(values["rear_fog_two_lamps_present"], two_lamps)
				self.assertEqual(values["rear_fog_absent"], absent)

	def test_rear_fog_width_requirement_applies_only_to_one_lamp(self):
		protocol = SimpleNamespace()
		measurement = SimpleNamespace()
		cases = (
			(None, "не применяется (в ТС отсутствуют задние противотуманные фонари)"),
			(0, "не применяется (в ТС отсутствуют задние противотуманные фонари)"),
			(1, "соответствует"),
			(2, "не применяется (в ТС имеется два задних противотуманных фонаря)"),
		)

		for count, expected in cases:
			with self.subTest(rear_fog_count=count):
				light = SimpleNamespace(rear_fog_count=count)
				values = build_dynamic_result_values(protocol, measurement, light)

				self.assertEqual(values["result_a_8_13_1_status"], expected)
				self.assertEqual(values["rear_fog_a_8_13_1_status"], expected)
				self.assertEqual(
					values["rear_fog_a_8_13_1_conclusion"],
					values["result_a_8_13_1_conclusion"],
				)

	def test_rear_fog_clause_8_13_2_splits_status_and_conclusion(self):
		absent = build_rear_fog_values(SimpleNamespace(rear_fog_count=0))
		self.assertEqual(
			absent["rear_fog_a_8_13_2_status"],
			"не применяется (в ТС отсутствуют задние противотуманные фонари)",
		)
		self.assertEqual(absent["rear_fog_a_8_13_2_conclusion"], "-")

		present = build_rear_fog_values(SimpleNamespace(
			rear_fog_count=1,
			rear_fog_upper_point_mm=500,
			rear_fog_lower_point_mm=400,
		))
		self.assertEqual(present["rear_fog_a_8_13_2_status"], "соответствует")
		self.assertIn("Результат:", present["rear_fog_a_8_13_2_conclusion"])
		self.assertIn("±", present["rear_fog_a_8_13_2_conclusion"])

	def test_rear_fog_8_13_2_conditions_and_values_for_one_and_two_lamps(self):
		def values(count):
			return build_rear_fog_values(SimpleNamespace(
				rear_fog_count=count,
				rear_fog_upper_point_mm=945,
				rear_fog_lower_point_mm=855,
			))

		one, two, none = values(1), values(2), values(0)

		self.assertEqual((one["a_8_13_2_one_lamp"], one["a_8_13_2_two_lamps"]), (True, False))
		self.assertEqual((two["a_8_13_2_one_lamp"], two["a_8_13_2_two_lamps"]), (False, True))
		self.assertEqual((none["a_8_13_2_one_lamp"], none["a_8_13_2_two_lamps"]), (False, False))

		# Правый фонарь повторяет значения левого.
		self.assertEqual(two["a_8_13_2_left_lower_point_value"], "855,00 мм")
		self.assertEqual(two["a_8_13_2_right_lower_point_value"], "855,00 мм")
		self.assertEqual(two["a_8_13_2_left_upper_point_value"], "945,00 мм")
		self.assertEqual(two["a_8_13_2_right_upper_point_value"], "945,00 мм")

		two_lamps_result = "\n".join((
			"Результат:",
			"Левая нижняя граница: 855,00 мм ± 0,29 мм",
			"Левая верхняя граница: 945,00 мм ± 0,29 мм",
			"Правая нижняя граница: 855,00 мм ± 0,29 мм",
			"Правая верхняя граница: 945,00 мм ± 0,29 мм",
		))
		self.assertIn(two_lamps_result, two["rear_fog_a_8_13_2_conclusion"])

		one_lamp_result = "\n".join((
			"Результат:",
			"Верхняя граница: 945,00 мм ± 0,29 мм",
			"Нижняя граница: 855,00 мм ± 0,29 мм",
		))
		self.assertIn(one_lamp_result, one["rear_fog_a_8_13_2_conclusion"])
		self.assertEqual(none["rear_fog_a_8_13_2_conclusion"], "-")

	def test_dynamic_non_applicability_statuses_use_excel_messages(self):
		values = build_dynamic_result_values(
			SimpleNamespace(tire_season="summer", has_spikes=False),
			SimpleNamespace(
				glonass_button_present=False,
				steering_lock_present=False,
				steps_present=False,
				opening_roof_present=False,
				spare_wheel_present=False,
				gas_equipment_present=False,
			),
			SimpleNamespace(
				front_fog_count=0,
				rear_fog_count=0,
				daytime_running_light_count=0,
				parking_light_count=0,
				rear_parking_light_count=0,
				adaptive_front_lighting_count=0,
				headlight_washer_present=False,
			),
		)

		self.assertEqual(
			values["result_a_3_2_status"],
			"не применяется (пункт Постановления Правительства)",
		)
		self.assertEqual(
			values["result_a_8_7_status"],
			"не применяется (в фарах ТС установлены источники света отличные от описанных в п. А.8.7)",
		)
		self.assertEqual(
			values["result_a_8_20_3_status"],
			"не применяется (ТС не оснащено устройствами фароочистки  и автоматическим корректирующим устройством угла наклона фар (не предусмотрено конструкцией))",
		)
		self.assertEqual(
			values["result_a_10_5_status"],
			"не применяется (на ТС установлены летние шины)",
		)

	def test_a87_and_a8203_depend_only_on_adaptive_front_lighting(self):
		protocol = SimpleNamespace()
		measurement = SimpleNamespace()
		cases = (
			(None, False),
			(0, True),
			(2, False),
			(2, True),
		)

		for adaptive_count, washer_present in cases:
			with self.subTest(
				adaptive_count=adaptive_count,
				washer_present=washer_present,
			):
				light = SimpleNamespace(
					adaptive_front_lighting_count=adaptive_count,
					headlight_washer_present=washer_present,
				)
				dynamic_values = build_dynamic_result_values(
					protocol,
					measurement,
					light,
				)
				applicability = build_applicability_values(
					protocol,
					measurement,
					light,
					dynamic_values,
				)

				expected = (
					"соответствует"
					if adaptive_count == 2
					else "не применяется (в фарах ТС установлены источники света "
					"отличные от описанных в п. А.8.7)"
				)
				expected_8_20_3 = (
					"соответствует"
					if adaptive_count == 2
					else "не применяется (ТС не оснащено устройствами фароочистки  и "
					"автоматическим корректирующим устройством угла наклона фар "
					"(не предусмотрено конструкцией))"
				)
				self.assertEqual(dynamic_values["result_a_8_7_status"], expected)
				self.assertEqual(applicability["applicable_8_7"], expected)
				self.assertEqual(
					applicability["applicable_8_20_3"],
					expected_8_20_3,
				)

	def test_common_applicability_statuses_match_vehicle_conditions(self):
		protocol = SimpleNamespace(tire_season="summer", has_spikes=False)
		measurement = SimpleNamespace(
			fuel_type="diesel",
			mileage_km=1000,
			glonass_button_present=False,
		)
		light = SimpleNamespace()
		dynamic_values = build_dynamic_result_values(protocol, measurement, light)
		values = build_applicability_values(protocol, measurement, light, dynamic_values)

		self.assertEqual(
			values["applicable_3_1"],
			"не применяется (пункт Постановления Правительства)",
		)
		self.assertEqual(
			values["applicable_21_7"],
			"не применяется (пробег ТС менее 3000 км)",
		)
		self.assertEqual(
			values["applicable_2_1"],
			"не применяется (ТС не используется для коммерческих перевозок)",
		)
		self.assertEqual(values["applicable_3_3_4"], "не указано")
		self.assertEqual(
			values["applicable_5_1_2_2"],
			"не применяется (ТС оборудован двухконтурной тормозной системой. "
			"Запасная тормозная система входит в состав рабочей тормозной системы "
			"и не оснащается независимым органом управления)",
		)
		self.assertEqual(
			values["applicable_8_19"],
			"не применяется (на ТС отсутствует светоотражающая маркировка)",
		)
		self.assertEqual(
			build_applicability_values(
				SimpleNamespace(vehicle_category="M1"),
				SimpleNamespace(fuel_type="diesel", vehicle_weight_kg=2000),
				SimpleNamespace(),
				{},
			)["applicable_21_3"],
			"не применяется (ТС категории М1 массой менее 3,5 т)",
		)
		for fuel_type, expected in (
			("petrol", "не применяется (ТС оборудовано только бензиновым двигателем)"),
			("hybrid_petrol", "не применяется (ТС оборудовано только бензиновым двигателем)"),
			("electric", "не применяется (ТС оборудовано только электродвигателем)"),
			("diesel", "не применяется (ТС категории М1 массой менее 3,5 т)"),
			("hybrid_diesel", "не применяется (ТС категории М1 массой менее 3,5 т)"),
		):
			with self.subTest(fuel_type=fuel_type):
				self.assertEqual(
					build_applicability_values(
						SimpleNamespace(vehicle_category="M1"),
						SimpleNamespace(fuel_type=fuel_type),
						SimpleNamespace(),
						{},
					)["applicable_21_3"],
					expected,
				)
		self.assertEqual(
			values["applicable_8_20_3"],
			"не применяется (ТС не оснащено устройствами фароочистки  и "
			"автоматическим корректирующим устройством угла наклона фар "
			"(не предусмотрено конструкцией))",
		)

		washer_present_light = SimpleNamespace(
			headlight_washer_present=True,
			adaptive_front_lighting_count=2,
		)
		washer_present_dynamic_values = build_dynamic_result_values(
			protocol,
			measurement,
			washer_present_light,
		)
		washer_present_values = build_applicability_values(
			protocol,
			measurement,
			washer_present_light,
			washer_present_dynamic_values,
		)
		self.assertEqual(washer_present_values["applicable_8_20_3"], "соответствует")
		for clause in ("8_18_1", "8_18_2", "8_18_3", "8_18_4"):
			self.assertEqual(values[f"applicable_{clause}"], "отсутствие")
	def test_a20_5_1_depends_only_on_electric_engine(self):
		cases = (
			(
				"electric",
				"fixed_cap",
				"не применяется (ТС оборудовано только электродвигателем)",
			),
			(
				"petrol",
				"fixed_cap",
				"не применяется (на ТС отсутствует несъемная крышка наливной горловины топливного бака)",
			),
			(
				"diesel",
				None,
				"не применяется (на ТС отсутствует несъемная крышка наливной горловины топливного бака)",
			),
		)

		for fuel_type, protection_measure, expected in cases:
			with self.subTest(fuel_type=fuel_type, protection_measure=protection_measure):
				values = build_applicability_values(
					SimpleNamespace(),
					SimpleNamespace(
						fuel_type=fuel_type,
						fuel_tank_leak_protection_measure=protection_measure,
					),
					SimpleNamespace(),
					{},
				)

				self.assertEqual(values["applicable_20_5_1"], expected)

	def test_a21_4_to_a21_6_apply_only_to_non_electric_engines(self):
		for fuel_type in ("electric", "petrol", "diesel"):
			with self.subTest(fuel_type=fuel_type):
				values = build_applicability_values(
					SimpleNamespace(),
					SimpleNamespace(fuel_type=fuel_type),
					SimpleNamespace(),
					{},
				)
				expected = (
					"не применяется (ТС оборудовано только электродвигателем)"
					if fuel_type == "electric"
					else "соответствует"
				)

				for clause in ("21_4", "21_5", "21_6"):
					self.assertEqual(values[f"applicable_{clause}"], expected)

	def test_a23_1_exposes_electric_and_numeric_docx_conditions(self):
		for fuel_type, is_electric in (
			("electric", True),
			("petrol", False),
			("diesel", False),
			("hybrid_petrol", False),
			("hybrid_diesel", False),
			(None, False),
		):
			with self.subTest(fuel_type=fuel_type):
				values = build_dynamic_result_values(
					SimpleNamespace(),
					SimpleNamespace(fuel_type=fuel_type),
					SimpleNamespace(),
				)

				self.assertEqual(values["a_23_1_electric"], is_electric)
				self.assertEqual(values["a_23_1_numeric"], not is_electric)

	def test_a23_2_has_separate_applicability_and_conclusion_placeholders(self):
		protocol = SimpleNamespace()
		light = SimpleNamespace()
		cases = (
			(
				"electric",
				"не применяется (ТС оборудовано только электродвигателем)",
				"-",
			),
			(
				"petrol",
				"соответствует",
				"Соответствует требованиям\nТР ТС 018/2011\nПриложения № 8 п.9.10",
			),
		)

		for fuel_type, expected_status, expected_conclusion in cases:
			with self.subTest(fuel_type=fuel_type):
				measurement = SimpleNamespace(fuel_type=fuel_type)
				context = build_dynamic_result_values(protocol, measurement, light)
				applicability = build_applicability_values(
					protocol,
					measurement,
					light,
					context,
				)

				self.assertEqual(context["result_a_23_2_status"], expected_status)
				self.assertEqual(
					context["result_a_23_2_conclusion"],
					expected_conclusion,
				)
				self.assertEqual(applicability["applicable_23_2"], expected_status)

	def test_a24_6_has_separate_applicability_and_conclusion_placeholders(self):
		cases = (
			(
				"electric",
				"не применяется (ТС оборудовано электрической рулевой рейкой)",
				"-",
			),
			(
				"hydraulic",
				"соответствует",
				"Соответствует требованиям\nТР ТС 018/2011\nПриложения № 8 п.2.6",
			),
		)

		for booster_type, expected_status, expected_conclusion in cases:
			with self.subTest(steering_booster_type=booster_type):
				protocol = SimpleNamespace()
				measurement = SimpleNamespace(steering_booster_type=booster_type)
				light = SimpleNamespace()
				context = build_dynamic_result_values(protocol, measurement, light)
				applicability = build_applicability_values(
					protocol,
					measurement,
					light,
					context,
				)

				self.assertEqual(context["result_a_24_6_status"], expected_status)
				self.assertEqual(
					context["result_a_24_6_conclusion"],
					expected_conclusion,
				)
				self.assertEqual(applicability["applicable_24_6"], expected_status)

	def test_a26_7_uses_pneumatic_suspension_presence(self):
		protocol = SimpleNamespace()
		cases = (
			(
				False,
				"не применяется (на ТС отсутствует пневматическая подвеска)",
				"-",
			),
			(
				True,
				"соответствует",
				"Соответствует требованиям\nТР ТС 018/2011\nПриложения № 8 п.10.7",
			),
		)

		for present, expected_status, expected_conclusion in cases:
			with self.subTest(pneumatic_suspension_present=present):
				measurement = SimpleNamespace(
					pneumatic_suspension_present=present,
				)
				context = build_dynamic_result_values(
					protocol,
					measurement,
					SimpleNamespace(),
				)
				applicability = build_applicability_values(
					protocol,
					measurement,
					SimpleNamespace(),
					context,
				)

				self.assertEqual(context["result_a_26_7_status"], expected_status)
				self.assertEqual(context["result_a_26_7_conclusion"], expected_conclusion)
				self.assertEqual(applicability["applicable_26_7"], expected_status)

	def test_a20_5_2_depends_only_on_electric_engine(self):
		cases = (
			(
				"electric",
				"structural_elements",
				"не применяется (ТС оборудовано только электродвигателем)",
			),
			(
				"petrol",
				"structural_elements",
				"не применяется (на ТС отсутствуют элементы конструкции, "
				"не допускающие утечки избыточных паров и топлива в случае "
				"отсутствия крышки наливной горловины)",
			),
			(
				"diesel",
				"fixed_cap",
				"не применяется (на ТС отсутствуют элементы конструкции, "
				"не допускающие утечки избыточных паров и топлива в случае "
				"отсутствия крышки наливной горловины)",
			),
		)

		for fuel_type, protection_measure, expected in cases:
			with self.subTest(fuel_type=fuel_type, protection_measure=protection_measure):
				values = build_applicability_values(
					SimpleNamespace(),
					SimpleNamespace(
						fuel_type=fuel_type,
						fuel_tank_leak_protection_measure=protection_measure,
					),
					SimpleNamespace(),
					{},
				)

				self.assertEqual(values["applicable_20_5_2"], expected)

	def test_full_result_fields_use_excel_reasons(self):
		self.assertEqual(
			build_front_fog_values(SimpleNamespace(front_fog_count=0))[
				"full_result_a_8_10_1"
			],
			"не применяется (в ТС отсутствуют передние противотуманные фары)",
		)
		self.assertEqual(
			build_rear_fog_values(SimpleNamespace(rear_fog_count=0))[
				"full_result_a_8_13_2"
			],
			"не применяется (в ТС отсутствуют задние противотуманные фонари)",
		)
		self.assertEqual(
			build_tire_depth_values(
				SimpleNamespace(tire_season="summer"),
				SimpleNamespace(),
			)["full_result_a_10_7_3"],
			"не применяется (на ТС установлены летние шины)",
		)


	def test_front_fog_absence_marks_all_three_clauses_not_applicable(self):
		expected = "не применяется (в ТС отсутствуют передние противотуманные фары)"

		for count in (None, 0):
			with self.subTest(front_fog_count=count):
				protocol = SimpleNamespace()
				measurement = SimpleNamespace()
				light = SimpleNamespace(front_fog_count=count)
				context = build_front_fog_values(light)
				context.update(build_dynamic_result_values(protocol, measurement, light))

				values = build_applicability_values(
					protocol,
					measurement,
					light,
					context,
				)

				for clause in ("8_10_1", "8_10_2", "8_10_3"):
					self.assertEqual(values[f"applicable_{clause}"], expected)

	def test_optional_lights_mark_all_screenshot_clauses_not_applicable(self):
		expected_front_fog = (
			"не применяется (в ТС отсутствуют передние противотуманные фары)"
		)
		expected_rear_fog = (
			"не применяется (в ТС отсутствуют задние противотуманные фонари)"
		)
		protocol = SimpleNamespace()
		measurement = SimpleNamespace()

		for absent_count in (None, 0):
			with self.subTest(absent_count=absent_count):
				light = SimpleNamespace(
					front_fog_count=absent_count,
					rear_fog_count=absent_count,
					turn_signal_count=absent_count,
				)
				context = build_front_fog_values(light)
				context.update(build_rear_fog_values(light))
				context.update(build_dynamic_result_values(protocol, measurement, light))
				values = build_applicability_values(
					protocol,
					measurement,
					light,
					context,
				)

				for clause in ("8_10_1", "8_10_2", "8_10_3"):
					self.assertEqual(values[f"applicable_{clause}"], expected_front_fog)
				self.assertEqual(values["applicable_8_13_1"], expected_rear_fog)
				self.assertEqual(values["applicable_8_13_2"], expected_rear_fog)

	def test_eco_checks_apply_fuel_rules_and_inclusive_mileage_threshold(self):
		protocol = SimpleNamespace()
		cases = [
			("diesel", 1000, "не применяется (пробег ТС менее 3000 км)", "не применяется (пробег ТС менее 3000 км)"),
			("electric", 5000, "не применяется (ТС оборудовано только электродвигателем)", "не применяется (ТС оборудовано только электродвигателем)"),
			("petrol", 2999, "не применяется (пробег ТС менее 3000 км)", "не применяется (пробег ТС менее 3000 км)"),
			("petrol", 3000, "соответствует", "не применяется (ТС оборудовано бензиновым двигателем)"),
			("diesel", 3000, "не применяется (ТС оборудовано дизельным двигателем)", "соответствует"),
			# Гибрид-бензин — как бензин, гибрид-дизель — как дизель.
			("hybrid_petrol", 3000, "соответствует", "не применяется (ТС оборудовано бензиновым двигателем)"),
			("hybrid_diesel", 3000, "не применяется (ТС оборудовано дизельным двигателем)", "соответствует"),
			(None, None, "соответствует", "соответствует"),
		]

		for fuel_type, mileage, co_status, smoke_status in cases:
			with self.subTest(fuel_type=fuel_type, mileage=mileage):
				measurement = SimpleNamespace(fuel_type=fuel_type, mileage_km=mileage)
				values = build_dynamic_result_values(
					protocol,
					measurement,
					SimpleNamespace(),
				)
				eco_values = build_eco_values(protocol, measurement)

				self.assertEqual(values["result_a_21_7_status"], co_status)
				self.assertEqual(values["result_a_21_8_status"], smoke_status)
				self.assertEqual(
					eco_values["a_21_7_low_mileage"],
					mileage is not None and mileage < 3000,
				)
				self.assertEqual(
					eco_values["a_21_7_numeric"],
					co_status == "соответствует",
				)
				self.assertEqual(
					eco_values["a_21_8_low_mileage"],
					mileage is not None and mileage < 3000,
				)
				self.assertEqual(
					eco_values["a_21_8_numeric"],
					smoke_status == "соответствует",
				)
				self.assertEqual(
					eco_values["mileage_21_9"],
					"не применяется (ТС оборудовано только электродвигателем)"
					if fuel_type == "electric"
					else "более 3000 км"
					if mileage is None or mileage >= 3000
					else "менее 3000 км",
				)
				if fuel_type is None and mileage is None:
					self.assertNotEqual(eco_values["full_result_a_21_9"], "-")
				self.assertEqual(
					values["result_a_21_9_status"],
					"не применяется (ТС оборудовано только электродвигателем)"
					if fuel_type == "electric"
					else "более 3000 км"
					if mileage is None or mileage >= 3000
					else "менее 3000 км",
				)


class ProtocolNumericRangeTests(SimpleTestCase):
	def test_condition_serializer_reports_label_and_bounds(self):
		serializer = ProtocolTestConditionSerializer(
			data={"ambient_temperature_c": "21.1"},
			partial=True,
		)

		self.assertFalse(serializer.is_valid())
		message = str(serializer.errors["ambient_temperature_c"][0])
		self.assertIn("Температура окружающей среды", message)
		self.assertIn("от 19 до 21", message)
		self.assertIn("°C", message)

	def test_explicit_dash_or_empty_value_skips_range_check(self):
		serializer = ProtocolTestConditionSerializer(
			data={
				"ambient_temperature_c": None,
				"dash_fields": ["ambient_temperature_c"],
			},
			partial=True,
		)

		self.assertTrue(serializer.is_valid(), serializer.errors)

	def measurement_with_axles(self, axle1, axle2):
		# 700 кг + 4 места (2/2) * 75 = 1000 кг → допустимо 950–990 кг.
		return ProtocolMeasurementSerializer(
			data={
				"vehicle_weight_kg": "700",
				"seats_count": "2/2",
				"axle1_load_kg": axle1,
				"axle2_load_kg": axle2,
			},
			partial=True,
		)

	def test_axis_sum_within_one_to_five_percent_below_full_mass_is_valid(self):
		for axle2 in ("450", "470", "490"):
			with self.subTest(axle2=axle2):
				serializer = self.measurement_with_axles("500", axle2)
				self.assertTrue(serializer.is_valid(), serializer.errors)

	def test_axis_sum_outside_one_to_five_percent_is_rejected(self):
		for axle2 in ("449", "491", "600"):
			with self.subTest(axle2=axle2):
				serializer = self.measurement_with_axles("500", axle2)
				self.assertFalse(serializer.is_valid())
				message = str(serializer.errors["axle1_load_kg"][0])
				self.assertIn("950–990 кг", message)
				self.assertIn("массы Drom с учетом мест", message)

	def test_measurer_save_skips_drom_mass_range_check(self):
		data = {"vehicle_weight_kg": "700", "seats_count": "2/2", "axle1_load_kg": "500", "axle2_load_kg": "600"}

		measurer = ProtocolMeasurementSerializer(data=data, partial=True, context={"skip_drom_mass_check": True})
		self.assertTrue(measurer.is_valid(), measurer.errors)

		operator = ProtocolMeasurementSerializer(data=data, partial=True)
		self.assertFalse(operator.is_valid())
		self.assertIn("axle1_load_kg", operator.errors)

	def test_fuel_type_from_catalog_splits_hybrids(self):
		normalize = ProtocolCreateSerializer().normalize_fuel_type
		self.assertEqual(normalize("Гибрид"), "hybrid_petrol")
		self.assertEqual(normalize("гибрид (бензин)"), "hybrid_petrol")
		self.assertEqual(normalize("Гибрид дизель"), "hybrid_diesel")
		self.assertEqual(normalize("Бензин"), "petrol")
		self.assertEqual(normalize("Дизель"), "diesel")

	def test_mass_range_is_checked_only_when_axles_are_in_request(self):
		measurement = SimpleNamespace(
			protocol=SimpleNamespace(configuration_id=None, tire_season=None),
			vehicle_weight_kg=Decimal("700"),
			seats_count="2/2",
			axle1_load_kg=Decimal("500"),
			axle2_load_kg=Decimal("600"),
		)

		without_axles = ProtocolMeasurementSerializer(
			instance=measurement,
			data={"seats_count": "2/2", "engine_model": "3ZR-FAE"},
			partial=True,
		)
		self.assertTrue(without_axles.is_valid(), without_axles.errors)

		with_axles = ProtocolMeasurementSerializer(
			instance=measurement,
			data={"axle2_load_kg": "600"},
			partial=True,
		)
		self.assertFalse(with_axles.is_valid())
		self.assertIn("950–990 кг", str(with_axles.errors["axle1_load_kg"][0]))

	def test_mass_range_shows_fractional_bounds(self):
		low, high = calc_axle_mass_range_kg(Decimal("1450"))
		self.assertEqual((format_mass_kg(low), format_mass_kg(high)), ("1377,5", "1435,5"))

	def test_missing_drom_mass_or_seats_skips_axis_sum_comparison(self):
		for data in (
			{"vehicle_weight_kg": None, "seats_count": "2/3"},
			{"vehicle_weight_kg": "1000", "seats_count": None},
		):
			with self.subTest(data=data):
				serializer = ProtocolMeasurementSerializer(
					data={**data, "axle1_load_kg": "1740", "axle2_load_kg": "1500"},
					partial=True,
				)

				self.assertTrue(serializer.is_valid(), serializer.errors)

	def test_seats_are_summed_by_rows(self):
		self.assertEqual(count_seats("2/3"), 5)
		self.assertEqual(count_seats("2/2/3"), 7)
		self.assertEqual(count_seats(" 2 / 3 "), 5)
		self.assertEqual(count_seats("2"), 2)
		self.assertIsNone(count_seats(""))
		self.assertIsNone(count_seats(None))
		self.assertIsNone(count_seats("5 мест"))
		self.assertEqual(calc_drom_full_mass_kg("1000", "2/2/2"), Decimal("1450"))
		self.assertIsNone(calc_drom_full_mass_kg(None, "2/3"))

	def test_brake_mass_comes_only_from_axle_loads(self):
		self.assertEqual(
			get_vehicle_mass_for_brake_calc(
				SimpleNamespace(axle1_load_kg="800", axle2_load_kg="600", vehicle_weight_kg="1000")
			),
			Decimal("1400"),
		)
		self.assertIsNone(
			get_vehicle_mass_for_brake_calc(
				SimpleNamespace(axle1_load_kg=None, axle2_load_kg=None, vehicle_weight_kg="1000")
			)
		)

	def test_power_supply_uses_excel_voltage_limits(self):
		serializer = ProtocolPowerSupplySerializer(
			data={"phase_a_n_voltage_v": "186.99"},
			partial=True,
		)

		self.assertFalse(serializer.is_valid())
		self.assertIn("phase_a_n_voltage_v", serializer.errors)
		self.assertIn("187", str(serializer.errors["phase_a_n_voltage_v"][0]))
		self.assertIn("242", str(serializer.errors["phase_a_n_voltage_v"][0]))

	def test_brake_difference_uses_fixed_fourteen_percent_limit(self):
		within_limit = ProtocolBrakeSerializer(
			data={
				"service_brake_type": "disc_disc",
				"axle_1_brake_difference_pct": "14",
				"axle_2_brake_difference_pct": "14",
			},
			partial=True,
		)
		above_limit = ProtocolBrakeSerializer(
			data={"service_brake_type": "disc_drum", "axle_1_brake_difference_pct": "14.01"},
			partial=True,
		)

		self.assertTrue(within_limit.is_valid(), within_limit.errors)
		self.assertFalse(above_limit.is_valid())
		self.assertIn("14", str(above_limit.errors["axle_1_brake_difference_pct"][0]))

	def test_speedometer_range_and_actual_speed_calculation(self):
		serializer = ProtocolMeasurementSerializer(
			data={"actual_speed_kmh": "999", "speed_by_speedometer_kmh": "20"},
			partial=True,
		)

		self.assertTrue(serializer.is_valid(), serializer.errors)
		self.assertEqual(serializer.validated_data["actual_speed_kmh"], Decimal("19.0"))

		accepted = ProtocolMeasurementSerializer(data={"speed_by_speedometer_kmh": "25.00"}, partial=True)
		self.assertTrue(accepted.is_valid(), accepted.errors)

		# По таблице диапазонов скорость по спидометру — только 20 или 25.
		for invalid_speed in ("19.99", "21", "25.01", "30"):
			with self.subTest(speedometer_speed=invalid_speed):
				invalid = ProtocolMeasurementSerializer(
					data={"speed_by_speedometer_kmh": invalid_speed},
					partial=True,
				)
				self.assertFalse(invalid.is_valid())
				self.assertIn("speed_by_speedometer_kmh", invalid.errors)

	def test_non_headlight_light_count_uses_allowed_discrete_values(self):
		serializer = ProtocolLightSerializer(
			data={"reverse_light_count": 3},
			partial=True,
		)

		self.assertFalse(serializer.is_valid())
		message = str(serializer.errors["reverse_light_count"][0])
		self.assertIn("1 или 2", message)
		self.assertIn("от 1 до 2", message)

	def test_turn_signal_frequency_accepts_all_available_values(self):
		serializer = ProtocolLightSerializer(
			data={"turn_signal_frequency_hz": "1.5", "turn_signal_frequency_per_min": "90"},
			partial=True,
		)

		self.assertTrue(serializer.is_valid(), serializer.errors)

	def test_lighting_measurement_ranges_match_reference_table(self):
		cases = (
			("low_beam_upper_point_mm", "500.99"),
			("low_beam_lower_point_mm", "1199.01"),
			("fog_light_left_distance_mm", "399.01"),
			("fog_light_upper_point_mm", "250.99"),
			("fog_light_upper_point_mm", "799.01"),
			("left_34v_cd", "99.99"),
			("left_34v_cd", "642.01"),
			("left_52h_cd", "1899.99"),
			("right_52h_cd", "110000.01"),
			("left_high_beam_cd", "4999.99"),
			("right_high_beam_cd", "112000.01"),
			("brake_signal_upper_point_mm", "1499.01"),
			("rear_fog_lower_point_mm", "250.99"),
		)

		for field_name, value in cases:
			with self.subTest(field_name=field_name, value=value):
				serializer = ProtocolLightSerializer(data={field_name: value}, partial=True)
				self.assertFalse(serializer.is_valid())
				self.assertIn(field_name, serializer.errors)

	def test_steering_backlash_and_glass_transparency_ranges_match_reference_table(self):
		cases = (
			(ProtocolMeasurementSerializer, "steering_backlash_deg", "9.51"),
			(ProtocolMeasurementSerializer, "glass_transparency_right_pct", "71.99"),
			(ProtocolMeasurementSerializer, "glass_transparency_left_pct", "95.01"),
			(ProtocolMeasurementSerializer, "glass_transparency_windshield_pct", "71.99"),
			(ProtocolMeasurementSerializer, "co_min_pct", "0.30"),
			(ProtocolMeasurementSerializer, "co_max_pct", "0.20"),
			(ProtocolMeasurementSerializer, "exhaust_noise_constant_db", "94.01"),
			(ProtocolMeasurementSerializer, "bumper_to_body_distance_mm", "19.01"),
			(ProtocolMeasurementSerializer, "vehicle_length_mm", "8148.01"),
		)

		for serializer_class, field_name, value in cases:
			with self.subTest(field_name=field_name, value=value):
				serializer = serializer_class(data={field_name: value}, partial=True)
				self.assertFalse(serializer.is_valid())
				self.assertIn(field_name, serializer.errors)


	def test_new_table_bounds_are_inclusive(self):
		cases = (
			(ProtocolTestConditionSerializer, {"ambient_temperature_c": "19", "relative_humidity_pct": "66", "atmospheric_pressure_kpa": "93"}),
			(ProtocolMeasurementSerializer, {"axle1_load_kg": "0", "axle2_load_kg": "3500", "co_min_pct": "0.29", "glass_transparency_right_pct": "72"}),
			(ProtocolLightSerializer, {"left_34v_cd": "100", "right_34v_cd": "642", "left_high_beam_cd": "5000", "fog_light_lower_point_mm": "251"}),
		)
		for serializer_class, data in cases:
			with self.subTest(data=data):
				serializer = serializer_class(data=data, partial=True)
				self.assertTrue(serializer.is_valid(), serializer.errors)

	def test_static_protocol_fields_are_not_range_checked(self):
		# В протоколе эти данные статичны — проверок по ним нет.
		self.assertTrue(ProtocolPowerSupplySerializer(data={"frequency_hz": "49"}, partial=True).is_valid())
		self.assertTrue(ProtocolMeasurementSerializer(data={"vehicle_weight_kg": "100"}, partial=True).is_valid())
		light = ProtocolLightSerializer(
			data={"turn_signal_count": 3, "low_beam_count": 1, "main_brake_signal_count": 3},
			partial=True,
		)
		self.assertTrue(light.is_valid(), light.errors)

	def test_optional_lights_allow_zero_or_two(self):
		for value, valid in ((0, True), (2, True), (1, False)):
			with self.subTest(value=value):
				serializer = ProtocolLightSerializer(
					data={"daytime_running_light_count": value, "front_fog_count": value},
					partial=True,
				)
				self.assertEqual(serializer.is_valid(), valid, serializer.errors)

	def test_parking_control_force_limits_depend_on_control_type(self):
		for parking_type, limit in (("mechanical_hand", "373"), ("mechanical_pedal", "465")):
			with self.subTest(parking_type=parking_type):
				ok = ProtocolBrakeSerializer(
					data={"parking_brake_type": parking_type, "parking_brake_control_force_n": limit},
					partial=True,
				)
				self.assertTrue(ok.is_valid(), ok.errors)
				too_big = ProtocolBrakeSerializer(
					data={"parking_brake_type": parking_type, "parking_brake_control_force_n": str(int(limit) + 1)},
					partial=True,
				)
				self.assertFalse(too_big.is_valid())


class ProtocolCalculationTests(TestCase):
	def test_calculated_vehicle_mass_is_the_sum_of_both_axles(self):
		self.assertEqual(
			calc_axle_load_sum_kg(
				SimpleNamespace(axle1_load_kg=Decimal("500"), axle2_load_kg=Decimal("340"))
			),
			Decimal("840"),
		)
		self.assertIsNone(calc_axle_load_sum_kg(SimpleNamespace(axle1_load_kg=500)))

	def test_light_absorption_uncertainty_uses_type_b_for_equal_readings(self):
		measurement = SimpleNamespace(
			light_absorption_1=Decimal("0.10"),
			light_absorption_2=Decimal("0.10"),
			light_absorption_3=Decimal("0.10"),
			light_absorption_4=Decimal("0.10"),
			light_absorption_5=Decimal("0.10"),
			light_absorption_6=Decimal("0.10"),
		)

		uncertainty = calc_light_absorption_uncertainty(measurement)

		self.assertAlmostEqual(float(uncertainty), 0.0476338, places=6)

	def test_light_absorption_uncertainty_increases_with_reading_spread(self):
		stable = SimpleNamespace(**{
			f"light_absorption_{index}": Decimal("0.10")
			for index in range(1, 7)
		})
		spread = SimpleNamespace(
			light_absorption_1=Decimal("0.00"),
			light_absorption_2=Decimal("0.02"),
			light_absorption_3=Decimal("0.10"),
			light_absorption_4=Decimal("0.12"),
			light_absorption_5=Decimal("0.20"),
			light_absorption_6=Decimal("0.22"),
		)

		self.assertGreater(
			calc_light_absorption_uncertainty(spread),
			calc_light_absorption_uncertainty(stable),
		)

	def test_headlight_uncertainty_includes_resolution_component(self):
		relative_only = Decimal("450") * Decimal("15") / Decimal("100") / Decimal("3").sqrt()
		updated = standard_headlight_uncertainty_cd(Decimal("450"))

		self.assertGreater(updated, relative_only)

	def test_control_force_uncertainty_includes_resolution_component(self):
		relative_only = Decimal("98") * Decimal("5") / Decimal("100") / Decimal("3").sqrt()

		self.assertGreater(
			u_control_force_n(Decimal("98")),
			Decimal("1.65") * relative_only,
		)

	def test_turn_signal_frequency_uncertainty_includes_resolution_component(self):
		old_relative_only = Decimal("1.65") * Decimal("0.1") / Decimal("3").sqrt()

		self.assertGreater(u_turn_signal_frequency_hz(), old_relative_only)

	def test_measurement_uncertainties_include_resolution_components(self):
		old_noise = Decimal("1.65") * Decimal("0.5") / Decimal("3").sqrt()
		old_glass = Decimal("1.65") * Decimal("2") / Decimal("3").sqrt()
		old_speed = Decimal("1.65") * Decimal("20") * Decimal("0.15") / Decimal("100") / Decimal("3").sqrt()

		self.assertGreater(u_noise_db(), old_noise)
		self.assertGreater(u_glass_transparency_pct(), old_glass)
		self.assertGreater(u_speed_kmh(Decimal("20")), old_speed)

	def test_axle_load_uncertainty_uses_brake_stand(self):
		# Стенд: ±3 % и цена деления 1 кг; 995 кг -> uB = 17,2363 (в прототипе без цены деления 17,2339), U = 28,44.
		self.assertAlmostEqual(float(u_axle_load_kg(Decimal("995")) / Decimal("1.65")), 17.2363, places=3)
		self.assertEqual(
			build_uncertainty_values(SimpleNamespace(measurement=SimpleNamespace(axle1_load_kg=Decimal("900"))))["u_axle1_load_kg"],
			"25,73",
		)
		self.assertIsNone(u_axle_load_kg(None))

	def test_steering_and_height_uncertainties_use_new_instrument_components(self):
		old_steering = Decimal("1.65") * Decimal("0.5") / Decimal("3").sqrt()
		old_height = Decimal("1.65") * Decimal("0.3") / Decimal("3").sqrt()

		self.assertGreater(u_steering_backlash_deg(), old_steering)
		self.assertNotEqual(u_vehicle_height_mm(1640), old_height)


class ProtocolDocxConditionalRowTests(SimpleTestCase):
	def test_table_row_else_renders_only_the_matching_branch(self):
		for condition, expected in ((True, "present"), (False, "absent")):
			with self.subTest(condition=condition):
				document = Document()
			table = document.add_table(rows=5, cols=1)
			table.cell(0, 0).text = "{%tr if front_fog_present %}"
			table.cell(1, 0).text = "present"
			table.cell(2, 0).text = "{%tr else %}"
			table.cell(3, 0).text = "absent"
			table.cell(4, 0).text = "{%tr endif %}"

			process_document_conditions(
				document,
				{"front_fog_present": condition},
			)

			self.assertEqual(
				[row.cells[0].text for row in document.tables[0].rows],
				[expected],
			)

	def test_table_row_elif_renders_first_matching_branch(self):
		cases = (
			({"one": True, "two": False}, ["один", "после"]),
			({"one": False, "two": True}, ["два", "после"]),
			({"one": False, "two": False}, ["нет", "после"]),
			# Если истинны оба, срабатывает только первая ветка.
			({"one": True, "two": True}, ["один", "после"]),
		)
		for context, expected in cases:
			with self.subTest(context=context):
				document = Document()
				table = document.add_table(rows=8, cols=1)
				for index, text in enumerate((
					"{%tr if one %}", "один",
					"{%tr elif two %}", "два",
					"{%tr else %}", "нет",
					"{%tr endif %}", "после",
				)):
					table.cell(index, 0).text = text

				process_document_conditions(document, context)

				self.assertEqual([row.cells[0].text for row in document.tables[0].rows], expected)
