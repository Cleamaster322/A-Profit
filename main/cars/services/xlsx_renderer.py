"""
Заполнение Excel-шаблона (openpyxl) по тем же правилам, что и DOCX:

- {{ key }} в ячейке заменяется значением из контекста. Если в ячейке только один
  плейсхолдер, а значение — число, в ячейку пишется число (формат ячейки из шаблона,
  например «0,00», сохраняется);
- строки с {%tr if … %} / {%tr elif … %} / {%tr else %} / {%tr endif %} в любой ячейке
  управляют выводом строк ниже и сами удаляются, как в renderer.process_table_conditions.

Объединенные ячейки допустимы только выше первой условной строки: openpyxl не сдвигает
объединения при удалении строк.
"""
import math
import re

from openpyxl import load_workbook

from .protocol_docx.renderer import (
    PLACEHOLDER_RE,
    TR_ELIF_RE,
    TR_ELSE_RE,
    TR_ENDIF_RE,
    TR_IF_RE,
    evaluate_tr_condition,
)

SINGLE_PLACEHOLDER_RE = re.compile(r"^\s*{{\s*([a-zA-Z0-9_]+)\s*}}\s*$")
LINE_HEIGHT_PT = 15


def row_text(worksheet, row_index):
    return " ".join(
        str(cell.value)
        for cell in worksheet[row_index]
        if cell.value is not None
    )


def rows_to_delete(worksheet, context):
    """Индексы строк, которые нужно удалить: служебные {%tr %} и строки ложных веток."""
    condition_stack = []
    result = []

    for row_index in range(1, worksheet.max_row + 1):
        text = row_text(worksheet, row_index)
        if_match = TR_IF_RE.search(text)
        elif_match = TR_ELIF_RE.search(text)
        else_match = TR_ELSE_RE.search(text)
        endif_match = TR_ENDIF_RE.search(text)

        if if_match:
            parent_active = all(frame["active"] for frame in condition_stack)
            condition_result = evaluate_tr_condition(if_match.group(1), context)
            condition_stack.append({
                "parent_active": parent_active,
                "condition_result": condition_result,
                "active": parent_active and condition_result,
                "else_seen": False,
            })
            result.append(row_index)
            continue

        if elif_match:
            result.append(row_index)
            if condition_stack and not condition_stack[-1]["else_seen"]:
                frame = condition_stack[-1]
                branch_result = (
                    not frame["condition_result"]
                    and evaluate_tr_condition(elif_match.group(1), context)
                )
                frame["active"] = frame["parent_active"] and branch_result
                frame["condition_result"] = frame["condition_result"] or branch_result
            continue

        if else_match:
            result.append(row_index)
            if condition_stack and not condition_stack[-1]["else_seen"]:
                frame = condition_stack[-1]
                frame["active"] = frame["parent_active"] and not frame["condition_result"]
                frame["else_seen"] = True
            continue

        if endif_match:
            if condition_stack:
                condition_stack.pop()
            result.append(row_index)
            continue

        if condition_stack and not all(frame["active"] for frame in condition_stack):
            result.append(row_index)

    return result


def fill_placeholders(worksheet, context):
    for row in worksheet.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or "{{" not in cell.value:
                continue

            single = SINGLE_PLACEHOLDER_RE.match(cell.value)
            if single:
                value = context.get(single.group(1), "")
                cell.value = value if isinstance(value, (int, float)) else ("" if value is None else str(value))
                continue

            cell.value = PLACEHOLDER_RE.sub(
                lambda match: "" if context.get(match.group(1)) is None else str(context.get(match.group(1))),
                cell.value,
            )


def fit_row_heights(worksheet):
    """
    Высота строк с переносом текста: openpyxl сам не подбирает её, а после удаления
    строк высоты из шаблона «съезжают». Считаем по длине текста и ширине колонки.
    Строки с объединенными ячейками (шапка) не трогаем — их высота задана в шаблоне.
    """
    widths = {
        letter: (dimension.width or 8.43)
        for letter, dimension in worksheet.column_dimensions.items()
    }
    merged_rows = {
        row_index
        for merged in worksheet.merged_cells.ranges
        for row_index in range(merged.min_row, merged.max_row + 1)
    }

    for row_index in range(1, worksheet.max_row + 1):
        if row_index in merged_rows:
            continue
        lines = 1
        for cell in worksheet[row_index]:
            if cell.value is None:
                continue
            width = widths.get(cell.column_letter, 8.43)
            chars_per_line = max(int(width * 1.1), 1)
            text_lines = str(cell.value).split("\n")
            cell_lines = sum(max(1, math.ceil(len(line) / chars_per_line)) for line in text_lines)
            lines = max(lines, cell_lines)
        worksheet.row_dimensions[row_index].height = LINE_HEIGHT_PT * lines


def render_xlsx(template_path, output_path, context):
    workbook = load_workbook(template_path)
    worksheet = workbook.active

    for row_index in sorted(rows_to_delete(worksheet, context), reverse=True):
        worksheet.delete_rows(row_index)

    fill_placeholders(worksheet, context)
    fit_row_heights(worksheet)

    workbook.save(output_path)
    return output_path
