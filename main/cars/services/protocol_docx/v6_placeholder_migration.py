import re
from functools import lru_cache
from pathlib import Path

from docx import Document

from .renderer import build_run_ranges, context_value_to_bool, find_run_index, iter_paragraphs


MAPPING_PATH = Path(__file__).with_name("V5_PLACEHOLDER_RENAME_MAP.md")
TEMPLATE_DIR = Path(__file__).resolve().parents[2] / "templates"
SOURCE_TEMPLATE = TEMPLATE_DIR / "protocol_template_v5_source.docx"
V6_TEMPLATE = TEMPLATE_DIR / "protocol_template_v6_source.docx"
EXPECTED_V5_PLACEHOLDER_COUNT = 253

PLACEHOLDER_LINE_RE = re.compile(
    r"^\s*{{\s*([A-Za-z0-9_]+)\s*}}\s*->\s*{{\s*([A-Za-z0-9_]+)\s*}}\s*$"
)
CONDITION_LINE_RE = re.compile(
    r"^\s*({%\s*tr\s+if\s+(?:not\s+)?[A-Za-z0-9_]+\s*%})"
    r"\s*->\s*"
    r"({%\s*tr\s+if\s+(?:not\s+)?[A-Za-z0-9_]+\s*%})\s*$"
)
TOKEN_RE = re.compile(
    r"{{\s*[A-Za-z0-9_]+\s*}}"
    r"|{%\s*tr\s+(?:if\s+(?:not\s+)?[A-Za-z0-9_]+|else|endif)\s*%}"
)
PLACEHOLDER_TOKEN_RE = re.compile(r"{{\s*([A-Za-z0-9_]+)\s*}}")
CONDITION_TOKEN_RE = re.compile(r"{%\s*tr\s+(if\s+(?:not\s+)?[A-Za-z0-9_]+)\s*%}")
CONDITION_EXPRESSION_RE = re.compile(r"^(if\s+)?(not\s+)?([A-Za-z0-9_]+)$")


@lru_cache(maxsize=1)
def load_v5_v6_mappings():
    if not MAPPING_PATH.exists():
        raise FileNotFoundError(f"Не найдена карта переименования: {MAPPING_PATH}")

    placeholder_renames = {}
    condition_renames = {}

    for line in MAPPING_PATH.read_text(encoding="utf-8").splitlines():
        placeholder_match = PLACEHOLDER_LINE_RE.match(line)
        if placeholder_match:
            old_name, new_name = placeholder_match.groups()
            previous = placeholder_renames.get(old_name)
            if previous is not None and previous != new_name:
                raise ValueError(f"Для {old_name} задано несколько новых имён")
            placeholder_renames[old_name] = new_name
            continue

        condition_match = CONDITION_LINE_RE.match(line)
        if condition_match:
            old_tag, new_tag = condition_match.groups()
            condition_renames[" ".join(old_tag.split())] = " ".join(new_tag.split())

    if len(placeholder_renames) != EXPECTED_V5_PLACEHOLDER_COUNT:
        raise ValueError(
            f"В карте найдено {len(placeholder_renames)} placeholders; "
            f"ожидалось {EXPECTED_V5_PLACEHOLDER_COUNT}"
        )

    return placeholder_renames, condition_renames


def add_v6_context_aliases(context):
    placeholder_renames, condition_renames = load_v5_v6_mappings()
    missing_keys = set(placeholder_renames) - set(context)
    if missing_keys:
        raise KeyError(
            "В DOCX-контексте отсутствуют legacy keys: "
            + ", ".join(sorted(missing_keys))
        )

    aliases = {}
    for old_name, new_name in placeholder_renames.items():
        aliases[new_name] = context[old_name]

    # The legacy applicable field contains a composite result. Prefer the
    # dedicated conclusion field, which correctly becomes '-' when not applicable.
    conclusion_source = "rear_fog_a_8_13_2_conclusion"
    if conclusion_source in context:
        aliases["a_8_13_2_conclusion"] = context[conclusion_source]

    for old_tag, new_tag in condition_renames.items():
        old_match = CONDITION_TOKEN_RE.fullmatch(old_tag)
        new_match = CONDITION_TOKEN_RE.fullmatch(new_tag)
        if not old_match or not new_match:
            continue

        old_expression = CONDITION_EXPRESSION_RE.match(old_match.group(1))
        new_expression = CONDITION_EXPRESSION_RE.match(new_match.group(1))
        if not old_expression or not new_expression:
            continue

        old_negated, old_name = old_expression.group(2), old_expression.group(3)
        new_negated, new_name = new_expression.group(2), new_expression.group(3)
        if old_name not in context:
            raise KeyError(f"В DOCX-контексте отсутствует условие {old_name}")

        condition_value = context_value_to_bool(context[old_name])
        if old_negated:
            condition_value = not condition_value
        aliases[new_name] = not condition_value if new_negated else condition_value

    context.update(aliases)
    return context


def _replace_paragraph_tokens(paragraph, placeholder_renames, condition_renames):
    if not paragraph.runs:
        return

    full_text = "".join(run.text for run in paragraph.runs)
    matches = list(TOKEN_RE.finditer(full_text))

    for match in reversed(matches):
        token = match.group(0)
        placeholder_match = PLACEHOLDER_TOKEN_RE.fullmatch(token)
        if placeholder_match:
            new_name = placeholder_renames.get(placeholder_match.group(1))
            if new_name is None:
                continue
            replacement = "{{ " + new_name + " }}"
        else:
            normalized_token = " ".join(token.split())
            replacement = condition_renames.get(normalized_token)
            if replacement is None:
                continue

        ranges = build_run_ranges(paragraph)
        start_index = find_run_index(ranges, match.start())
        end_index = find_run_index(ranges, match.end() - 1)
        if start_index is None or end_index is None:
            continue

        first_run, first_start, _ = ranges[start_index]
        last_run, last_start, _ = ranges[end_index]
        before = first_run.text[:match.start() - first_start]
        after = last_run.text[match.end() - last_start:]

        if start_index == end_index:
            first_run.text = before + replacement + after
            continue

        first_run.text = before + replacement
        for run_index in range(start_index + 1, end_index):
            paragraph.runs[run_index].text = ""
        last_run.text = after


def create_v6_template(destination=None):
    if not SOURCE_TEMPLATE.exists():
        raise FileNotFoundError(f"Не найден исходный шаблон v5: {SOURCE_TEMPLATE}")

    output_path = Path(destination) if destination else V6_TEMPLATE
    if output_path.exists():
        raise FileExistsError(f"Шаблон v6 уже существует: {output_path}")

    placeholder_renames, condition_renames = load_v5_v6_mappings()
    document = Document(SOURCE_TEMPLATE)

    source_conditions = {
        " ".join(match.group().split())
        for paragraph in iter_paragraphs(document)
        for match in CONDITION_TOKEN_RE.finditer(paragraph.text)
    }
    unmapped_conditions = source_conditions - set(condition_renames)
    if unmapped_conditions:
        raise ValueError(
            "В карте отсутствуют row conditions: "
            + ", ".join(sorted(unmapped_conditions))
        )

    for paragraph in iter_paragraphs(document):
        _replace_paragraph_tokens(paragraph, placeholder_renames, condition_renames)

    remaining_conditions = {
        " ".join(match.group().split())
        for paragraph in iter_paragraphs(document)
        for match in CONDITION_TOKEN_RE.finditer(paragraph.text)
    } & source_conditions
    if remaining_conditions:
        raise ValueError(
            "В шаблоне v6 остались старые row conditions: "
            + ", ".join(sorted(remaining_conditions))
        )

    rendered_placeholders = {
        match.group(1)
        for paragraph in iter_paragraphs(document)
        for match in PLACEHOLDER_TOKEN_RE.finditer(paragraph.text)
    }
    expected_placeholders = set(placeholder_renames.values())
    if rendered_placeholders != expected_placeholders:
        raise ValueError(
            "Набор placeholders v6 не совпадает с картой переименования: "
            f"missing={sorted(expected_placeholders - rendered_placeholders)}, "
            f"extra={sorted(rendered_placeholders - expected_placeholders)}"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
